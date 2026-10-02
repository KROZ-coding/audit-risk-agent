"""产物级一致性回归：PDF/Excel 与独立预期字面量逐项比对。

背景（方案「B：产物级一致性核对」）：以往导出测试只验证「能生成、不崩溃」，
未验证产物里的数字/期间/口径/单位是否与真实计算工具输出一致；风险证据原文
也未做机器字段泄漏（judgment_1/overall_status/risk_chain_analysis 等内部
结构化键）的回归。本文件把中国石油 2025 半年报真实 fixture 跑一遍
validate → calculate → 导出，再用独立写死的预期字面量核对 PDF 提取文本与
Excel 单元格，防止「工具算对了、导出层数字漂移」这类缺陷复发。

测试原则：
- 预期字面量独立于被测函数（不从被测代码反推数字）；
- 金额/期间/口径/单位/指标 5 类关键信息在 PDF 与 Excel 中必须同时出现；
- 机器字段/内部键（judgment_1、overall_status、risk_chain_analysis、
  ###、---）一律不得泄漏到客户可见产物；
- 负例：量级错误（总资产×100）必须被数据校验拦截，不得进入报告。
"""
import json
import os
import re
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))

from offline_real_verify import CALC_JSON, VD_JSON

from tools.data_validator import validate_financial_data
from tools.excel_export import _export_excel_impl
from tools.financial_calculator import calculate_financial_indicators
from tools.pdf_export import _export_pdf_impl


def _invoke(tool, data):
    return json.loads(tool.invoke({"financial_data_json": json.dumps(data, ensure_ascii=False)}))


def _run_tools():
    """跑真实确定性工具，返回 (fin, vd) 完整结果。"""
    fin = _invoke(calculate_financial_indicators, CALC_JSON)
    vd = _invoke(validate_financial_data, VD_JSON)
    return fin, vd


def _base_report(fin):
    """构造与真实工具输出同源的报告台账（导出层会实时重算 risk_summary）。"""
    return {
        "company_info": {
            "company_name": "中国石油天然气股份有限公司",
            "stock_code": "601857",
            "report_year": "2025（半年度）",
            "industry": "能源",
            "audit_opinion": "未经审计（半年度报告）",
        },
        "risk_summary": {"total_risks": 1, "major_risks": 1, "important_risks": 0,
                         "general_risks": 0, "risk_dimensions": {}},
        "risk_details": [
            {
                "risk_id": "R001",
                "dimension": "财务错报风险",
                "title": "应收账款增速高于营收，占营业收入8.26%",
                "level": "重大",
                "confidence": 0.85,
                "evidence": "2025年6月30日应收账款119,715百万元，较上年末71,610百万元。"
                           "关联方提供产品和服务占同类交易14.96%。"
                           "数据期间2025年半年度，口径中国准则合并，金额单位人民币元。",
                "data_analysis": "应收增速显著高于营收增速，需核查信用政策与收入确认。",
                "regulatory_basis": "《企业会计准则第14号——收入》",
                "case_reference": "某上市公司应收账款异常增长案例",
                "audit_suggestion": "对主要欠款方执行函证程序并核查账龄。",
            }
        ],
        "overall_assessment": "综合风险评分 未获取/无法判定（请人工复核）。",
        "comprehensive_score_snapshot": {"score": None, "level": "未获取", "level_key": "unknown"},
        "review_gate": {"status": "not_run"},
        "metric_results": fin.get("metric_results"),
        "result_schema_version": "1.0",
        "rule_version": "2026-09-v3",
        "period": "2025年半年度",
        "scope": "中国准则合并",
        "amount_unit": "人民币元",
    }


# ── 独立预期字面量（不与被测函数反推；折行敏感处用短子串） ──
PDF_EXPECTED = [
    "中国石油天然气股份有限公司", "601857", "2025（半年度）", "未经审计（半年度报告）",
    "2025年半年度", "中国准则合并", "人民币元",
    "119,715百万元", "71,610百万元", "8.26%", "14.96",
    "综合风险评分 未获取/无法判定",
]

XLSX_EXPECTED = [
    # 封面与契约
    "中国石油天然气股份有限公司", "601857", "报告年度", "2025（半年度）",
    "数据期间", "2025年半年度", "数据口径", "中国准则合并",
    "风险摘要", "系统采信风险", "重要风险",
    "结果契约版本", "1.0", "计算版本", "校验版本",
    # 底稿章节标题
    "原始事实", "指标计算", "勾稽校验", "证据索引", "待复核事项",
    # 金额/单位/指标（独立字面量）
    "1,450,099,000,000.00人民币元", "119,715,000,000.00人民币元",
    "20.89%", "38.48%", "8.26%", "14.96%",
    "181.00天（按已声明期间日历计算）",
]

# 机器字段/内部键：不得泄漏到客户可见 PDF/Excel
LEAKAGE_KEYS = ["judgment_1", "overall_status", "risk_chain_analysis", "###", "---",
                "E-SOURCE-", "E-SYSTEM-", "E-RISK-MODEL-"]


@pytest.fixture(autouse=True)
def _stub_upload(monkeypatch):
    """屏蔽真实落盘到 local_storage/（两个导出函数内部才 import 该模块属性）。"""
    monkeypatch.setattr(
        "local_storage.upload_file_to_storage",
        lambda *a, **k: "/local_storage/reports/stubbed",
    )


def _pdf_text(path):
    import pypdf
    reader = pypdf.PdfReader(str(path))
    return "\n".join((pg.extract_text() or "") for pg in reader.pages)


def _xlsx_text(path):
    from openpyxl import load_workbook
    wb = load_workbook(str(path), data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(ws.title)  # sheet 名也是底稿结构的一部分
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if v is not None:
                    parts.append(str(v))
    return "\n".join(parts)


class TestPdfArtifactConsistency:
    """PDF 产物：封面/证据/结论与真实工具输出及期间口径一致，机器字段不泄漏。"""

    def test_legacy_single_pdf_contains_literal_facts_and_no_leakage(self, tmp_path):
        fin, _ = _run_tools()
        out = tmp_path / "r.pdf"
        result = _export_pdf_impl(
            json.dumps(_base_report(fin), ensure_ascii=False), str(out),
            financial_indicators_json=json.dumps(fin, ensure_ascii=False),
        )
        assert "已生成" in result and out.exists() and out.stat().st_size > 0
        text = _pdf_text(out)
        for literal in PDF_EXPECTED:
            assert literal in text, f"PDF 缺少独立预期字面量: {literal!r}"
        for key in LEAKAGE_KEYS:
            assert key not in text, f"PDF 泄漏内部机器字段: {key!r}"
        assert "财务风险" in text
        assert "财务错报风险" not in text
        assert "财务错报" not in text

    def test_split_summary_pdf_uses_same_period_scope_and_marks_score_unknown(self, monkeypatch, tmp_path):
        """拆分模式（默认路径）综合汇总：期间/口径/评分未获取口径与独立预期一致。"""
        from tools import pdf_export as pe
        fin, _ = _run_tools()
        saved = []
        monkeypatch.setattr("local_storage.upload_file_to_storage",
                            lambda path, dest, mime: (saved.append(path), f"/local_storage/{dest}")[1])
        monkeypatch.setattr(pe, "_embed_chart", lambda *a, **k: None)
        result = _export_pdf_impl(
            json.dumps(_base_report(fin), ensure_ascii=False),
            financial_indicators_json=json.dumps(fin, ensure_ascii=False),
        )
        assert result.count("已生成") == 3
        synth = [p for p in saved if p.endswith(".pdf") and "综合汇总报告" in os.path.basename(p)][0]
        text = "\n".join((pg.extract_text() or "") for pg in __import__("pypdf").PdfReader(synth).pages)
        assert "2025年半年度" in text
        assert "中国准则合并" in text
        assert "综合风险评分 未获取/无法判定" in text
        assert "119,715百万元" in text
        assert "8.26%" in text
        for key in LEAKAGE_KEYS:
            assert key not in text, f"综合汇总 PDF 泄漏内部机器字段: {key!r}"
        assert "财务风险" in text
        assert "财务错报风险" not in text
        assert "财务错报" not in text

    def test_pdf_cover_exposes_snapshot_release_context(self, tmp_path):
        report = _base_report({})
        report["company_info"].update({
            "accounting_standard": "中国企业会计准则",
            "scope": "合并",
            "amount_unit": "人民币元",
            "source_document": "sample.pdf",
            "source_file_sha256": "a" * 64,
        })
        report["source"] = {"document_name": "sample.pdf", "source_hash": "a" * 64}
        report["review_gate"] = {"status": "not_passed", "human_review_required": True}
        out = tmp_path / "release-context.pdf"
        result = _export_pdf_impl(json.dumps(report, ensure_ascii=False), str(out))
        assert "已生成" in result and out.exists()
        text = _pdf_text(out)
        for literal in ("中国准则合并", "人民币元", "not_passed", "正式采信风险", "待复核提示",
                        "人工复核"):
            assert literal in text, f"PDF 封面缺少发布上下文字段: {literal!r}"
        assert "a" * 64 in re.sub(r"\s+", "", text)


class TestExcelArtifactConsistency:
    """Excel 底稿：9 张表结构与金额/期间/口径/指标字面量一致，机器字段不泄漏。"""

    def test_excel_contains_independent_literals_and_no_leakage(self, tmp_path):
        fin, vd = _run_tools()
        out = tmp_path / "r.xlsx"
        result = _export_excel_impl(
            json.dumps(_base_report(fin), ensure_ascii=False), str(out),
            financial_indicators_json=json.dumps(fin, ensure_ascii=False),
            validation_json=json.dumps(vd, ensure_ascii=False),
        )
        assert "已生成" in result and out.exists() and out.stat().st_size > 0
        text = _xlsx_text(out)
        for literal in XLSX_EXPECTED:
            assert literal in text, f"Excel 缺少独立预期字面量: {literal!r}"
        assert "财务风险" in text
        assert "财务错报风险" not in text
        assert "财务错报" not in text
        for key in LEAKAGE_KEYS:
            assert key not in text, f"Excel 泄漏内部机器字段: {key!r}"

    def test_excel_converts_internal_page_markers_to_readable_citations(self, tmp_path):
        report = _base_report({})
        report["facts"] = [{
            "fact_id": "F-PAGE",
            "field": "source_excerpt",
            "raw_value": "--- 第 1 页 ---\\n营业收入原文",
            "value": "--- 第 1 页 ---\\n营业收入原文",
            "source_document": "sample.pdf",
            "page": 1,
        }]
        out = tmp_path / "page-marker.xlsx"
        result = _export_excel_impl(json.dumps(report, ensure_ascii=False), str(out))
        assert "已生成" in result and out.exists()
        text = _xlsx_text(out)
        assert "第 1 页" in text
        assert "---" not in text

    def test_excel_hides_internal_evidence_keys_but_keeps_readable_reference(self, tmp_path):
        report = _base_report({})
        report["risk_details"][0]["evidence_ids"] = [
            "E-SOURCE-R001-P85", "E-SYSTEM-ALERT-1", "E-RISK-MODEL-Z",
        ]
        report["evidence"] = [{
            "evidence_id": "E-SOURCE-R001-P85", "source_type": "source_report",
            "source_document": "sample.pdf", "page": "85", "verified": True,
            "status": "verified", "excerpt": "应收账款原文",
        }]
        out = tmp_path / "evidence-keys.xlsx"
        result = _export_excel_impl(json.dumps(report, ensure_ascii=False), str(out))
        assert "已生成" in result and out.exists()
        text = _xlsx_text(out)
        for key in ("E-SOURCE-", "E-SYSTEM-", "E-RISK-MODEL-"):
            assert key not in text, f"Excel 泄漏内部证据键: {key!r}"
        assert "源报告第 85 页" in text
        assert "系统计算证据" in text
        assert "量化模型证据" in text


class TestMagnitudeErrorGate:
    """负例：量级错误必须在数据校验环节被拦截，不得进入报告。"""

    def test_validator_rejects_100x_asset_magnitude_error(self):
        bad = dict(VD_JSON)
        bad["total_assets_current"] = VD_JSON["total_assets_current"] * 100
        out = _invoke(validate_financial_data, bad)
        dv = out.get("data_validation") or {}
        assert dv.get("validation_result") == "未通过"
        checks = {c.get("check"): c for c in dv.get("all_checks", [])}
        assert checks["资产负债表平衡"]["passed"] is False

    def test_validator_rejects_liability_equals_assets_swap(self):
        bad = dict(VD_JSON)
        bad["total_liabilities_current"] = VD_JSON["total_assets_current"]
        out = _invoke(validate_financial_data, bad)
        dv = out.get("data_validation") or {}
        assert dv.get("validation_result") == "未通过"
        balance = next(c for c in dv.get("all_checks", []) if c.get("check") == "资产负债表平衡")
        assert balance["passed"] is False
        assert "超过内部2%筛查阈值" in balance.get("message", "")
