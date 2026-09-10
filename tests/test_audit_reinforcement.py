"""审计补强五项（涉及科目/适用认定/核查程序/所需材料/企业改进建议）回归测试。

整改要求：每条相关风险须同时给出企业改进建议、涉及科目、适用认定、核查程序及
所需材料；函证、监盘、控制测试、截止测试等一律写成待执行建议，不得宣称已经实施；
模型缺项时由系统模板补齐并标注来源，避免"模型没给"被静默呈现为空白交付物。
"""
import json

from tools.audit_reinforcement import (
    REINFORCEMENT_FIELDS,
    apply_reinforcement,
    ensure_reinforcement,
    get_reinforcement,
    reinforcement_cell,
    reinforcement_items,
)

_LABELS = [label for _key, label in REINFORCEMENT_FIELDS]


def _risk(**overrides) -> dict:
    """财务错报维度、标题含「应收账款」的条目（同时命中维度级与科目级模板）。"""
    risk = {
        "risk_id": "R001",
        "dimension": "财务错报风险",
        "title": "应收账款增速显著高于营业收入增速",
        "level": "重要",
        "confidence": 0.8,
        "evidence": "应收账款同比+67%，营业收入同比+5%。",
        "data_analysis": "回款周期拉长，坏账风险上升。",
        "audit_suggestion": "扩大函证范围",
    }
    risk.update(overrides)
    return risk


def _squeeze(text: str) -> str:
    """去除全部空白，抵消 PDF 抽取时分词插入的换行与空格。"""
    return "".join(str(text).split())


class TestTemplateFallback:
    """模型未给字段时按维度/科目确定性补齐，并如实标注来源。"""

    def test_all_five_fields_filled_when_missing(self):
        values = get_reinforcement(_risk())
        assert list(values) == [key for key, _label in REINFORCEMENT_FIELDS]
        assert all(values[key] for key in values)

    def test_procedures_are_pending_only(self):
        """核查程序只能写"待执行"，不得出现已实施口吻。"""
        values = get_reinforcement(_risk(involved_accounts="应收账款；存货"))
        assert all(p.endswith("（待执行）") for p in values["audit_procedures"])

    def test_source_recorded_as_template(self):
        risk = _risk()
        assert ensure_reinforcement(risk) is True
        assert risk["reinforcement_source"] == "template"

    def test_items_follow_declared_order(self):
        items = reinforcement_items(_risk())
        assert [label for label, _values in items] == _LABELS

    def test_cell_joins_values_with_separator(self):
        cell = reinforcement_cell(_risk(), "involved_accounts")
        assert cell == "应收账款；营业收入"


class TestModelPriority:
    """模型已有输出必须保留，模板只补空缺。"""

    def test_complete_model_output_kept(self):
        risk = _risk(
            involved_accounts=["应收账款"],
            assertions=["存在", "计价与分摊"],
            audit_procedures=["对主要客户执行函证（待执行）"],
            required_materials=["账龄分析表"],
            improvement_suggestions=["加强回款考核"],
        )
        assert ensure_reinforcement(risk) is True
        assert risk["reinforcement_source"] == "llm"
        assert risk["assertions"] == ["存在", "计价与分摊"]
        assert risk["audit_procedures"] == ["对主要客户执行函证（待执行）"]

    def test_partial_model_output_marked_as_merged(self):
        risk = _risk(involved_accounts=["应收账款"])
        ensure_reinforcement(risk)
        assert risk["reinforcement_source"] == "llm+template"
        assert risk["involved_accounts"] == ["应收账款"]
        assert risk["assertions"] and risk["required_materials"]

    def test_get_reinforcement_does_not_mutate_input(self):
        risk = _risk()
        get_reinforcement(risk)
        assert "involved_accounts" not in risk
        assert "reinforcement_source" not in risk

    def test_second_run_is_noop(self):
        risk = _risk()
        assert ensure_reinforcement(risk) is True
        assert ensure_reinforcement(risk) is False

    def test_template_provenance_not_upgraded_on_rerun(self):
        """重跑不得把模板补齐的内容改标成模型输出（否则 PDF 模板提示会被抹掉）。"""
        risk = _risk()
        ensure_reinforcement(risk)
        assert risk["reinforcement_source"] == "template"
        risk["title"] = "应收账款增速显著高于营业收入增速"  # 其他字段被外部改写
        ensure_reinforcement(risk)
        assert risk["reinforcement_source"] == "template"


class TestNormalization:
    """LLM 输出的字符串/列表/字典/多行/项目符号统一归一为去重列表。"""

    def test_multiline_and_bullets_deduped(self):
        values = get_reinforcement(_risk(
            audit_procedures="- 取得明细表\n- 取得明细表\n- 复核计算\n"))
        assert values["audit_procedures"] == ["取得明细表", "复核计算"]

    def test_inline_split_only_for_enum_fields(self):
        """科目与认定按顿号拆分，长句程序描述不拆。"""
        values = get_reinforcement(_risk(
            involved_accounts="应收账款、存货，货币资金",
            audit_procedures="程序一、程序二"))
        assert values["involved_accounts"] == ["应收账款", "存货", "货币资金"]
        assert values["audit_procedures"] == ["程序一、程序二"]

    def test_dict_and_semicolon_flattened(self):
        values = get_reinforcement(_risk(
            required_materials={"1": "明细账", "2": "原始凭证"},
            improvement_suggestions="完善核算流程；建立复核机制"))
        assert values["required_materials"] == ["明细账", "原始凭证"]
        assert values["improvement_suggestions"] == ["完善核算流程", "建立复核机制"]

    def test_blank_values_treated_as_missing(self):
        values = get_reinforcement(_risk(assertions="   ", audit_procedures=[]))
        assert values["assertions"] == ["存在", "完整性", "准确性", "计价与分摊"]
        assert values["audit_procedures"]


class TestSubjectDerivation:
    """科目关键词触发专项程序与材料；维度不明时不猜测认定。"""

    def test_receivable_hits_confirmation_and_aging(self):
        values = get_reinforcement(_risk())
        assert "应收账款" in values["involved_accounts"]
        assert any("函证" in p for p in values["audit_procedures"])
        assert any("账龄" in m for m in values["required_materials"])

    def test_unknown_dimension_defers_to_manual_review(self):
        risk = {"risk_id": "R009", "dimension": "市场风险",
                "title": "国际油价大幅波动", "evidence": "报告期内油价显著波动。"}
        values = get_reinforcement(risk)
        assert values["involved_accounts"] == ["待人工确认"]
        assert values["assertions"] == ["待人工确认"]
        assert all(p.endswith("（待执行）") for p in values["audit_procedures"])


class TestApplyReinforcement:
    """台账级写回：网页/PDF/Excel 同源。"""

    def test_writes_back_all_fields_and_source(self):
        report = {"risk_details": [_risk(), _risk(risk_id="R002", dimension="关联交易风险")]}
        assert apply_reinforcement(report) == 2
        for risk in report["risk_details"]:
            for key, _label in REINFORCEMENT_FIELDS:
                assert risk[key]
            assert risk["reinforcement_source"] in ("llm", "llm+template", "template")

    def test_second_apply_changes_nothing(self):
        report = {"risk_details": [_risk()]}
        apply_reinforcement(report)
        assert apply_reinforcement(report) == 0

    def test_tolerates_missing_or_broken_details(self):
        assert apply_reinforcement({}) == 0
        assert apply_reinforcement({"risk_details": "not-a-list"}) == 0


class TestPdfRendering:
    """PDF 明细中补强五项编号固定为 (6)-(10)，模板来源附加提示行。"""

    def _export(self, monkeypatch, report):
        saved = []

        def _fake(path, dest, mime):
            saved.append(path)
            return f"/local_storage/{dest}"

        monkeypatch.setattr("local_storage.upload_file_to_storage", _fake)
        monkeypatch.setattr("tools.pdf_export._embed_chart", lambda *a, **k: None)
        from tools.pdf_export import _export_pdf_impl
        result = _export_pdf_impl(json.dumps(report, ensure_ascii=False))
        assert result.count("已生成") == 3
        return saved

    def _financial_text(self, saved):
        import pypdf
        path = [p for p in saved if "_财务健康诊断报告_" in p][0]
        return _squeeze("".join((pg.extract_text() or "") for pg in pypdf.PdfReader(path).pages))

    def test_five_items_numbered_6_to_10(self, monkeypatch):
        report = {"company_info": {"company_name": "补强测试公司", "report_year": "2025"},
                  "risk_details": [_risk()]}
        apply_reinforcement(report)
        text = self._financial_text(self._export(monkeypatch, report))
        for idx, label in enumerate(_LABELS, 6):
            assert f"({idx}){label}：" in text

    def test_template_sourced_items_carry_notice(self, monkeypatch):
        report = {"company_info": {"company_name": "补强测试公司", "report_year": "2025"},
                  "risk_details": [_risk()]}
        apply_reinforcement(report)
        text = self._financial_text(self._export(monkeypatch, report))
        assert "（上述补强项由系统模板按风险维度生成" in text

    def test_model_sourced_items_have_no_notice(self, monkeypatch):
        report = {"company_info": {"company_name": "补强测试公司", "report_year": "2025"},
                  "risk_details": [_risk(
                      involved_accounts=["应收账款"], assertions=["存在"],
                      audit_procedures=["执行函证（待执行）"],
                      required_materials=["账龄分析表"],
                      improvement_suggestions=["加强回款考核"])]}
        apply_reinforcement(report)
        text = self._financial_text(self._export(monkeypatch, report))
        assert "（上述补强项由系统模板按风险维度生成" not in text


class TestExcelRendering:
    """Excel 台账与建议表同步新增五列，缺审计建议的条目不再整条漏出。"""

    def _export(self, tmp_path, report):
        from openpyxl import load_workbook

        from tools.excel_export import _export_excel_impl
        out = tmp_path / "r.xlsx"
        _export_excel_impl(json.dumps(report, ensure_ascii=False), str(out))
        return load_workbook(str(out))

    def _report(self) -> dict:
        return {"company_info": {"company_name": "补强测试公司", "report_year": "2025"},
                "risk_details": [_risk(audit_suggestion=""),
                                  _risk(risk_id="R002", dimension="信息披露合规风险",
                                        title="重大事项延迟披露", audit_suggestion="核查披露时点")]}

    def test_ledger_appends_five_columns(self, tmp_path):
        ws = self._export(tmp_path, self._report())["风险台账"]
        headers = [c.value for c in ws[1]]
        assert headers[-5:] == _LABELS
        assert ws.max_column == 17

    def test_suggestion_sheet_appends_five_columns(self, tmp_path):
        ws = self._export(tmp_path, self._report())["建议与核查程序"]
        headers = [c.value for c in ws[1]]
        assert headers[-5:] == _LABELS
        assert ws.max_column == 9

    def test_entry_without_suggestion_still_listed(self, tmp_path):
        ws = self._export(tmp_path, self._report())["建议与核查程序"]
        rows = {str(r[0]): r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
        assert set(rows) == {"R001", "R002"}
        assert rows["R001"][4]  # 涉及科目：模型未给，模板补齐后不得为空
        assert rows["R001"][6]  # 核查程序（待执行）
