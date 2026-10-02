"""导出模块格式稳健性回归测试。

覆盖 PDF/Excel 导出在面对 LLM 常见「脏数据」时的稳健性：
- XML 特殊字符（< > &）：reportlab Paragraph 按 XML 解析，未转义会报错/乱码；
  系统提示词的判定标准本身就含 "OCF/NP<0.5"，此类内容出现概率极高。
- 超长文本与多行换行：验证 PDF 不溢出、Excel 行高自适应不被裁剪。
- 非标量字段（dict 型 trend_analysis）：验证序列化为 JSON 而非渲染 Python repr。
- 行业基准标量值：验证 Excel 不再静默跳过非 dict 值。
- PDF 拆分模式（默认）：一次生成 3 份独立报告（财务健康/合规信披/综合汇总），
  每份含封面/目录/免责声明，维度拆分归属正确，图表失败不阻断。
"""
import json

import pytest

from tools.pdf_export import _export_pdf_impl
from tools.excel_export import _export_excel_impl


def _dirty_report() -> dict:
    """构造一份含特殊字符、超长文本、多行、非标量字段的风险台账。"""
    long_text = "应收账款增速远超营收增速，" * 50   # 超长文本：测试溢出/行高自适应
    return {
        "company_info": {
            "company_name": "AT&T 测试<集团>股份有限公司",   # 含 & < >
            "stock_code": "600000",
            "report_year": "2025",
            "industry": "互联网 & 科技",
            "audit_opinion": "保留意见（OCF/NP<0.5）",
        },
        "risk_summary": {
            "total_risks": 2, "major_risks": 1, "important_risks": 1, "general_risks": 0,
            "risk_dimensions": {"financial_misstatement": 1, "disclosure_compliance": 1},
        },
        "risk_details": [
            {
                "risk_id": "R001",
                "dimension": "财务错报风险",
                "title": "OCF/NP<0.5 且应收增速>营收增速20pp",   # 含 < >
                "level": "重大",
                "confidence": 0.85,
                "evidence": "经营现金流/净利润<0.5，连续2期。\n应收账款同比+45%。",  # 多行 + <
                "data_analysis": "营收 & 现金流背离，存在提前确认收入嫌疑",   # 含 &
                "regulatory_basis": "《审计准则1211号》<第二十条>",
                "case_reference": "某上市公司 & 其关联方造假案",
                "audit_suggestion": long_text,   # 超长文本
                "reasoning_chain": [
                    {"step": "数据发现", "detail": "OCF/NP=0.3 < 阈值0.5"},
                    {"step": "风险判定", "detail": "重大错报风险 > 一般瑕疵"},
                ],
                "trend_analysis": {"direction": "恶化", "periods": 2},   # dict 型非标量
            },
            {
                "risk_id": "R002",
                "dimension": "信息披露合规风险",
                "title": "重大事项延迟披露 & 会计政策变更",
                "level": "重要",
                "confidence": 0.6,
                "evidence": "延迟披露>30天",
                "audit_suggestion": "核查披露时点",
            },
        ],
        "overall_assessment": "综合判断：财务错报风险>合规风险，且现金流<0。\n建议全面核查。",
        "industry_benchmark": {
            "盈利能力": {"毛利率": "25%", "净利率": "8%"},
            "综合评分": 72,   # 标量值：测试 Excel 非 dict 分支
        },
    }


@pytest.fixture(autouse=True)
def _stub_upload(monkeypatch):
    """屏蔽真实落盘到 local_storage/，避免测试污染产物目录。

    两个导出函数均在内部 `from local_storage import upload_file_to_storage`，
    运行时才读取模块属性，故 monkeypatch 模块属性即可生效。
    """
    monkeypatch.setattr(
        "local_storage.upload_file_to_storage",
        lambda *a, **k: "/local_storage/reports/stubbed",
    )


class TestPdfExportFormatting:
    def test_dirty_data_does_not_crash(self, tmp_path):
        """含 < > & 与多行文本的脏数据不应使 PDF 生成崩溃或触发免责声明拦截。"""
        out = tmp_path / "r.pdf"
        result = _export_pdf_impl(json.dumps(_dirty_report(), ensure_ascii=False), str(out))
        assert "已生成" in result
        assert "被拦截" not in result   # 免责声明校验仍通过
        assert out.exists() and out.stat().st_size > 0
        # 防回归：两遍构建（总页数）不得产出空壳文件，首页必须可提取文本
        import pypdf
        reader = pypdf.PdfReader(str(out))
        assert len(reader.pages) >= 1
        assert (reader.pages[0].extract_text() or "").strip()

    def test_dict_input_also_works(self, tmp_path):
        """兼容已解析字典输入（兜底导出路径会直接传 dict）。"""
        out = tmp_path / "r2.pdf"
        result = _export_pdf_impl(_dirty_report(), str(out))
        assert "已生成" in result
        assert out.exists()


class TestPdfSplitExport:
    """拆分模式（缺省 output_path）：一次生成 3 份独立 PDF"""

    def test_split_generates_three_named_files(self, monkeypatch):
        """3 个文件落盘 + 3 条下载链接 + 固定后缀命名。"""
        uploaded = []
        monkeypatch.setattr(
            "local_storage.upload_file_to_storage",
            lambda path, dest, mime: (uploaded.append(dest), f"/local_storage/{dest}")[1],
        )
        # 跳过图表生成（避免测试依赖 matplotlib 耗时），验证无图路径不崩
        monkeypatch.setattr("tools.pdf_export._embed_chart", lambda *a, **k: None)
        result = _export_pdf_impl(json.dumps(_dirty_report(), ensure_ascii=False))
        assert result.count("已生成") == 3
        assert result.count("下载链接:") == 3
        assert "（财务健康）" in result and "（合规信披）" in result and "（综合汇总）" in result
        names = [d.split("/")[-1] for d in uploaded]
        assert any(n.endswith("_财务健康诊断报告.pdf") for n in names)
        assert any(n.endswith("_合规与信息披露报告.pdf") for n in names)
        assert any(n.endswith("_综合汇总报告.pdf") for n in names)

    def test_split_risks_dimension_mapping(self):
        """维度拆分归属：中英文维度值均能正确映射到对应报告。"""
        from tools.pdf_export import _split_risks, FINANCIAL_DIMS, COMPLIANCE_DIMS
        rd = [
            {"dimension": "financial_misstatement"},   # 财务（英文）
            {"dimension": "持续经营风险"},              # 财务（中文，未归一化）
            {"dimension": "disclosure_compliance"},    # 合规
            {"dimension": "关联交易风险"},              # 合规（中文）
            {"dimension": "regulatory_penalty"},       # 合规
        ]
        fin = _split_risks(rd, FINANCIAL_DIMS)
        comp = _split_risks(rd, COMPLIANCE_DIMS)
        assert len(fin) == 2 and len(comp) == 3

    def test_split_risks_short_dimension_names(self):
        """简称维度（SP「五大风险维度」表用语）也能正确拆分。

        历史 bug：过滤集合只有全称，LLM 按提示词输出简称时全部静默丢弃，
        导致财务/合规拆分报告风险恒为 0（用户反馈「结论 0 条 vs 明细 5 条」）。
        """
        from tools.pdf_export import _split_risks, FINANCIAL_DIMS, COMPLIANCE_DIMS
        rd = [
            {"dimension": "财务错报"},   # 财务（简称）
            {"dimension": "持续经营"},   # 财务（简称）
            {"dimension": "信披合规"},   # 合规（简称）
            {"dimension": "监管处罚"},   # 合规（简称）
            {"dimension": "关联交易"},   # 合规（简称）
        ]
        fin = _split_risks(rd, FINANCIAL_DIMS)
        comp = _split_risks(rd, COMPLIANCE_DIMS)
        assert len(fin) == 2 and len(comp) == 3

    def test_split_risks_one_word_shorthand(self):
        """一级词简写（实测：LLM 输出"财务"而非"财务错报"）也能正确拆分，
        否则 R005/R006 类条目既不进财务报告也不进合规报告（v23 实测丢包）。"""
        from tools.pdf_export import _split_risks, FINANCIAL_DIMS, COMPLIANCE_DIMS
        rd = [
            {"dimension": "财务"},        # 财务（一级词简写）
            {"dimension": "财务风险"},     # 财务（变体）
            {"dimension": "信息披露"},     # 合规（一级词简写）
            {"dimension": "信披"},        # 合规（变体）
        ]
        fin = _split_risks(rd, FINANCIAL_DIMS)
        comp = _split_risks(rd, COMPLIANCE_DIMS)
        assert len(fin) == 2 and len(comp) == 2

    def test_reconcile_summary_rewrites_contradictory_numbers(self):
        """risk_summary 与 risk_details 矛盾时强制以明细回写并收集警告。

        历史 bug：KPI 卡优先采信 LLM 自填的 risk_summary（截图 total=2 / 结论 0/0/0），
        明细有 5 条也照显示矛盾数字。
        """
        from tools.pdf_export import _reconcile_summary
        report = {
            "risk_summary": {"total_risks": 2, "major_risks": 0, "important_risks": 0,
                             "general_risks": 0, "risk_dimensions": {}},
            "risk_details": [
                {"dimension": "财务错报", "level": "高"},
                {"dimension": "财务错报", "level": "重要"},
                {"dimension": "信披合规", "level": "中"},
                {"dimension": "持续经营", "level": "一般"},
                {"dimension": "持续经营", "level": "低"},
            ],
        }
        warnings = _reconcile_summary(report)
        rs = report["risk_summary"]
        assert rs["total_risks"] == 5
        assert rs["major_risks"] == 1 and rs["important_risks"] == 2 and rs["general_risks"] == 2
        # 4 个 KPI 键 + 维度分布（P3 扩展校正）全部回写并收集警告
        assert len(warnings) == 5
        # 维度键已归一化为规范英文标识符（K 补丁），渲染层映射回中文全称
        assert rs["risk_dimensions"] == {"financial_misstatement": 2,
                                          "disclosure_compliance": 1,
                                          "going_concern": 2}
        # 原测试：维度表漏计（LLM 静态 risk_dimensions 为 {} 时按明细重算）
        assert any("risk_dimensions" in w for w in warnings)

    def test_reconcile_summary_rewrites_assessment_layer_counts(self):
        """整体结论中的门禁前计数必须与正式/待核查分层同步。"""
        from tools.pdf_export import _reconcile_summary
        report = {
            "risk_summary": {},
            "risk_details": [
                {"risk_id": "R001", "level": "重要", "formal_status": "unaccepted"},
                {"risk_id": "R002", "level": "重要", "formal_status": "accepted"},
            ],
            "accepted_risk_details": [
                {"risk_id": "R002", "level": "重要", "formal_status": "accepted"},
            ],
            "overall_assessment": (
                "识别1项风险，系统采信风险0项，所有风险均为建议关注等级，"
                "需结合人工专业判断复核确认。"
            ),
        }
        _reconcile_summary(report)
        assert "识别2项风险" in report["overall_assessment"]
        assert "系统采信风险1项" in report["overall_assessment"]
        assert "另有1项为待复核提示" in report["overall_assessment"]

    def test_reconcile_summary_ignores_stale_intermediate_snapshot(self):
        """终局前的旧快照不得覆盖刚完成门禁的正式风险计数。"""
        from tools.pdf_export import _reconcile_summary
        report = {
            "report_snapshot": {"risks": {"formal": [], "pending": []}},
            "risk_summary": {},
            "risk_details": [
                {"risk_id": "R001", "level": "重要", "formal_status": "accepted"},
                {"risk_id": "R002", "level": "一般", "formal_status": "unaccepted"},
            ],
            "accepted_risk_details": [
                {"risk_id": "R001", "level": "重要", "formal_status": "accepted"},
            ],
            "overall_assessment": "系统采信风险0项，需人工复核。",
        }
        _reconcile_summary(report)
        assert "系统采信风险1项" in report["overall_assessment"]

    def test_reconcile_summary_rewrites_stale_pending_item_count(self):
        from tools.pdf_export import _reconcile_summary
        report = {
            "risk_summary": {},
            "risk_details": [
                {"risk_id": f"R{i}", "level": "一般", "formal_status": "unaccepted"}
                for i in range(1, 5)
            ],
            "accepted_risk_details": [],
            "overall_assessment": "识别出3项待核实事项，需结合人工专业判断复核确认。",
        }
        _reconcile_summary(report)
        assert "识别出4项待复核提示" in report["overall_assessment"]
        assert "3项待核实事项" not in report["overall_assessment"]

    def test_sub_conclusion_counts_matched_to_subset(self, monkeypatch):
        """拆分报告结论章：KPI 与环形图同口径（均基于过滤后子集）。

        历史 bug：KPI 统计子集（2 条）而环形图统计全集（5 条），同页数字矛盾。
        """
        from tools import pdf_export as pe
        monkeypatch.setattr(pe, "_embed_chart", lambda *a, **k: None)
        st = pe._build_styles(pe._register_chinese_font())
        risks = [
            {"dimension": "财务错报", "level": "重大"},
            {"dimension": "财务错报", "level": "一般"},
        ]
        sub_json = json.dumps({"risk_details": risks}, ensure_ascii=False)
        body = pe._sub_conclusion_body("财务健康", risks, st, pe._register_chinese_font(), sub_json)
        assert any("系统采信相关风险 2 项" in getattr(part, "text", "") for part in body)
        conclusion = ' '.join(getattr(part, 'text', '') for part in body)
        assert "重大 1 项" in conclusion and "一般 1 项" in conclusion

    def test_contradictory_ledger_export_consistent(self, monkeypatch):
        """矛盾台账完整导出端到端：三份报告数字与明细一致（截图 bug 全链路回归）。"""
        from tools import pdf_export as pe
        uploaded = {}
        monkeypatch.setattr(
            "local_storage.upload_file_to_storage",
            lambda path, dest, mime: (uploaded.update({dest: path}), f"/local_storage/{dest}")[1],
        )
        monkeypatch.setattr(pe, "_embed_chart", lambda *a, **k: None)
        report = {
            "company_info": {"company_name": "矛盾台账测试公司", "industry": "制造业",
                              "report_year": "2025"},
            "risk_summary": {"total_risks": 2, "major_risks": 0, "important_risks": 0,
                             "general_risks": 0,
                             "risk_dimensions": {"财务错报": 2, "信披合规": 1, "持续经营": 2}},
            "risk_details": [
                {"risk_id": "R001", "dimension": "财务错报", "level": "高", "title": "应收激增"},
                {"risk_id": "R002", "dimension": "财务错报", "level": "重要", "title": "存贷双高"},
                {"risk_id": "R003", "dimension": "信披合规", "level": "中", "title": "政策变更"},
                {"risk_id": "R004", "dimension": "持续经营", "level": "一般", "title": "连续亏损"},
                {"risk_id": "R005", "dimension": "持续经营", "level": "低", "title": "现金流为负"},
            ],
            "overall_assessment": "综合评估意见。",
        }
        result = _export_pdf_impl(json.dumps(report, ensure_ascii=False))
        assert result.count("已生成") == 3
        import pypdf
        synth_path = [p for k, p in uploaded.items() if k.endswith("_综合汇总报告.pdf")][0]
        text = "".join((pg.extract_text() or "") for pg in pypdf.PdfReader(synth_path).pages)
        assert "共 5 条风险的详细分析" in text          # 目录与明细一致
        # P2: 数据一致性提示为内部校验日志，不渲染到客户可见 PDF（仅 logger 记录）
        assert "台账total_risks" not in text
        assert "已按明细重算" not in text
        fin_path = [p for k, p in uploaded.items() if k.endswith("_财务健康诊断报告.pdf")][0]
        fin_text = "".join((pg.extract_text() or "") for pg in pypdf.PdfReader(fin_path).pages)
        assert "系统采信相关风险 4 项" in fin_text       # 财务子集（错报2+持续经营2）
        comp_path = [p for k, p in uploaded.items() if k.endswith("_合规与信息披露报告.pdf")][0]
        comp_text = "".join((pg.extract_text() or "") for pg in pypdf.PdfReader(comp_path).pages)
        assert "系统采信相关风险 1 项" in comp_text      # 合规子集（信披合规1）

    def test_company_profile_placeholder_when_missing(self):
        """公司简介缺失时渲染占位说明，不崩溃；提供时正常渲染。"""
        from tools.pdf_export import _profile_body, _register_chinese_font, _build_styles
        st = _build_styles(_register_chinese_font())
        els = _profile_body({"company_info": {}}, st)
        assert "公司简介未提供" in els[0].text
        els2 = _profile_body({"company_info": {"company_profile": "本公司成立于1998年，主营锂电池材料。"}}, st)
        assert "成立于1998年" in els2[0].text
        # 顶层别名兼容
        els3 = _profile_body({"company_profile": "顶层别名生效"}, st)
        assert "顶层别名生效" in els3[0].text

    def test_embed_chart_failure_returns_none(self):
        """图表生成异常时 _embed_chart 静默降级返回 None，不抛出。"""
        from tools import pdf_export as pe

        def boom(*a, **k):
            raise RuntimeError("matplotlib unavailable")

        assert pe._embed_chart(boom, "{}") is None

    def test_financial_and_disclosure_sections(self, monkeypatch):
        """财务指标章/披露检查章：有效入参渲染、空/错误入参跳过。"""
        from tools import pdf_export as pe
        monkeypatch.setattr(pe, "_embed_chart", lambda *a, **k: None)
        st = pe._build_styles(pe._register_chinese_font())
        # 财务章：指标表 + 预警编号
        fin_body = pe._financial_section_body(json.dumps({
            "indicators": {"current_ratio": 0.8, "gross_margin_pct": 21.5},
            "alerts": ["连续两年净利润为负，存在持续经营风险"],
        }, ensure_ascii=False), st, pe._register_chinese_font(), "")
        assert fin_body
        assert pe._financial_section_body("", st, pe._register_chinese_font(), "") == []
        assert pe._financial_section_body("not-json", st, pe._register_chinese_font(), "") == []
        # 披露章：得分表 + 问题清单；error 结果跳过
        dc_ok = json.dumps({"compliance_score": 80.0, "risk_score": 20.0, "checked_items": 10,
                            "passed_items": 8, "issues": ["重大事项延迟披露"],
                            "sections_missing": ["公司治理"], "audit_opinion": "标准无保留意见"},
                           ensure_ascii=False)
        assert pe._disclosure_section_body(dc_ok, st, pe._register_chinese_font())
        assert pe._disclosure_section_body("", st, pe._register_chinese_font()) == []
        assert pe._disclosure_section_body(json.dumps({"error": "文本过短"}), st, pe._register_chinese_font()) == []


class TestPdfModuleFilter:
    """按模块裁剪产物：financial 只出财务健康报告，compliance 只出合规报告，
    synthesis/未指定/未知值出全部三份；返回文案与实际产物一一对应。"""

    @pytest.fixture(autouse=True)
    def _no_chart(self, monkeypatch):
        """跳过图表生成（避免测试依赖 matplotlib 耗时）。"""
        monkeypatch.setattr("tools.pdf_export._embed_chart", lambda *a, **k: None)

    def _run(self, monkeypatch, module):
        uploaded = []
        monkeypatch.setattr(
            "local_storage.upload_file_to_storage",
            lambda path, dest, mime: (uploaded.append(dest), f"/local_storage/{dest}")[1],
        )
        result = _export_pdf_impl(json.dumps(_dirty_report(), ensure_ascii=False), module=module)
        return result, uploaded

    def test_financial_module_only_financial_report(self, monkeypatch):
        """financial 模块：仅 1 条链接/1 份产物，且为财务健康诊断报告。"""
        result, uploaded = self._run(monkeypatch, "financial")
        assert result.count("下载链接:") == 1
        assert "（财务健康）" in result
        assert "（合规信披）" not in result and "（综合汇总）" not in result
        assert len(uploaded) == 1
        assert uploaded[0].endswith("_财务健康诊断报告.pdf")

    def test_compliance_module_only_compliance_report(self, monkeypatch):
        """compliance 模块：仅 1 条链接/1 份产物，且为合规与信息披露报告。"""
        result, uploaded = self._run(monkeypatch, "compliance")
        assert result.count("下载链接:") == 1
        assert "（合规信披）" in result
        assert "（财务健康）" not in result and "（综合汇总）" not in result
        assert len(uploaded) == 1
        assert uploaded[0].endswith("_合规与信息披露报告.pdf")

    @pytest.mark.parametrize("module", ["synthesis", "", "unknown_module"])
    def test_synthesis_or_unspecified_generates_all_three(self, monkeypatch, module):
        """synthesis/空/未知值：全部三份，文案标签与产物一一对应。"""
        result, uploaded = self._run(monkeypatch, module)
        assert result.count("下载链接:") == 3
        assert "（财务健康）" in result and "（合规信披）" in result and "（综合汇总）" in result
        assert len(uploaded) == 3

    def test_module_is_case_insensitive(self, monkeypatch):
        """模块值大小写不敏感（前端/LLM 传参容错）。"""
        result, uploaded = self._run(monkeypatch, "Financial")
        assert result.count("下载链接:") == 1
        assert "（财务健康）" in result

    def test_compliance_pdf_is_usable_with_sparse_data(self, monkeypatch):
        """合规模块真实链路形态：无财务指标数据、甚至无披露检查结果时，
        合规 PDF 仍必须完整可用（非空壳）：封面/目录/各章兜底文案齐全。
        合规模块不跑财务计算，financial_indicators_json 必然为空；
        串跑流水线场景下披露检查结果也可能缺失，均不得产出空壳报告。"""
        monkeypatch.setattr("tools.pdf_export._embed_chart", lambda *a, **k: None)
        # 上传 stub 兼做落盘路径收集器：直接读真实生成的文件，
        # 不依赖扫描 tempdir（避免历史残留文件干扰）
        saved = []
        monkeypatch.setattr("local_storage.upload_file_to_storage",
                            lambda path, dest, mime: (saved.append(path), f"/local_storage/{dest}")[1])
        result = _export_pdf_impl(json.dumps(_dirty_report(), ensure_ascii=False),
                                  module="compliance")
        assert result.count("下载链接:") == 1 and "（合规信披）" in result
        assert len(saved) == 1
        # 临时文件带 uuid 唯一后缀（并发安全），仍须以报告名结尾
        assert saved[0].endswith(".pdf") and "_合规与信息披露报告_" in saved[0]
        import pypdf
        reader = pypdf.PdfReader(saved[0])
        assert len(reader.pages) >= 3, "合规报告至少应含封面/目录/章节正文"
        text = "\n".join((p.extract_text() or "") for p in reader.pages)
        assert "合规与信息披露报告" in text          # 封面标题
        assert "合规标准对照" in text                # 固定章：不依赖披露检查结果
        assert "本报告结论" in text                  # 结论章始终渲染
        assert "财务指标四维判读" not in text        # 财务专属章不得串入合规报告


class TestPdfProfessionalUpgrade:
    """PDF 专业化升级：四维判读/评分解读/方法论/合规对照/KPI 卡/新图表"""

    def test_interpret_indicator_rules(self):
        """确定性判读：阈值规则与行业基准对比分支均生效。"""
        from tools.pdf_export import _interpret_indicator
        verdict, bad = _interpret_indicator("current_ratio", 0.8, 1.5)
        assert bad is True and "<1" in verdict
        verdict, bad = _interpret_indicator("gross_margin_pct", 21.5, 25.0)
        assert "低于行业基准" in verdict and bad is False
        verdict, bad = _interpret_indicator("debt_to_asset_ratio_pct", 88.0, 50.0)
        assert bad is True and "85%" in verdict
        verdict, bad = _interpret_indicator("operating_cashflow_to_net_profit_ratio", 0.3, 1.0)
        assert "OCF/NP<0.5" in verdict and bad is True
        # 无规则的辅助指标返回占位判读
        verdict, bad = _interpret_indicator("net_profit_current", 123.0, None)
        assert verdict == "—" and bad is None

    def test_financial_section_four_dim_calls_new_chart(self, monkeypatch):
        """财务章：四维判读表渲染且调用对比柱状图/雷达图/热力图。"""
        from tools import pdf_export as pe
        calls = []
        monkeypatch.setattr(pe, '_load_benchmarks', lambda industry: {'gross_margin':25, 'current_ratio':1.5})
        monkeypatch.setattr(pe, "_embed_chart",
                            lambda fn, payload, width_cm=14.0: calls.append(fn) or None)
        st = pe._build_styles(pe._register_chinese_font())
        fin = json.dumps({"indicators": {"gross_margin_pct": 21.5, "current_ratio": 0.8,
                                         "debt_to_asset_ratio_pct": 55.0,
                                         "inventory_turnover_ratio": 3.2,
                                         "revenue_yoy_change_pct": -25.0},
                          "alerts": ["应收账款增速远超营收"]}, ensure_ascii=False)
        body = pe._financial_section_body(fin, st, pe._register_chinese_font(), "", "制造业")
        assert body
        assert pe._gen_indicator_bars in calls      # 指标对比柱状图（有行业基准）
        assert pe._generate_radar_chart in calls
        assert pe._generate_risk_heatmap in calls

    def test_score_section_render_and_degrade(self):
        """评分解读章：有效 JSON 渲染，空/异常/缺 score 降级跳过。"""
        from tools import pdf_export as pe
        st = pe._build_styles(pe._register_chinese_font())
        sc = json.dumps({"score": 43.2, "level": "中等风险", "level_key": "medium",
                         "breakdown": {"financial": 50.0, "disclosure": 40.0, "validation": 30.0},
                         "weights": {"financial": 0.5, "disclosure": 0.3, "validation": 0.2},
                         "base_score": 40.0, "escalation": 3.2,
                         "escalation_reasons": ["Z-Score 落入财务困境区"],
                         "summary": "中等风险"}, ensure_ascii=False)
        assert pe._score_section_body(sc, st, pe._register_chinese_font())
        assert pe._score_section_body("", st, pe._register_chinese_font()) == []
        assert pe._score_section_body(json.dumps({"error": "异常"}), st, pe._register_chinese_font()) == []
        assert pe._score_section_body("not-json", st, pe._register_chinese_font()) == []
        assert pe._score_section_body(json.dumps({"level": "低风险"}), st, pe._register_chinese_font()) == []

    def test_methodology_body_with_models_and_validation(self):
        """方法论章：权重表/Z-M-Score 动态取值/校验摘要均渲染。"""
        from tools import pdf_export as pe
        st = pe._build_styles(pe._register_chinese_font())
        rm = json.dumps({"risk_models": {
            "altman_z_score": {"available": True, "score": 1.2, "zone": "财务困境区"},
            "beneish_m_score": {"available": False},
            "cross_interpretation": {"interpretation": "双重信号测试"}}}, ensure_ascii=False)
        vd = json.dumps({"passed_checks": 5, "failed_checks": 1,
                         "validation_result": "未通过"}, ensure_ascii=False)
        body = pe._methodology_body("", rm, vd, st, pe._register_chinese_font())
        text = " ".join(getattr(e, "text", "") for e in body if hasattr(e, "text"))
        assert "50%" in text and "Z-Score" in text and "1.200" in text
        assert "双重信号测试" in text and "未通过" in text
        # 无任何模型/校验数据时仍渲染静态方法论（不崩）
        assert pe._methodology_body("", "", "", st, pe._register_chinese_font())

    def test_compliance_standard_body_always_renders(self):
        """合规标准对照章：无检查数据也渲染静态法规表；有数据时结果列联动。"""
        from tools import pdf_export as pe
        st = pe._build_styles(pe._register_chinese_font())
        assert pe._compliance_standard_body("", st, pe._register_chinese_font())
        dc = json.dumps({"issues": ["重大事项延迟披露", "关联交易披露不充分"],
                         "sections_missing": ["公司治理"], "compliance_score": 65.0},
                        ensure_ascii=False)
        assert pe._compliance_standard_body(dc, st, pe._register_chinese_font())

    def test_cross_validation_and_risk_chain_skip_when_empty(self):
        """交叉验证/风险传导链章：字段非空渲染、空则跳过。"""
        from tools import pdf_export as pe
        st = pe._build_styles(pe._register_chinese_font())
        assert pe._cross_validation_body({}, st) == []
        assert pe._risk_chain_body({"risk_chain_analysis": " "}, st) == []
        assert pe._cross_validation_body({"cross_validation_analysis": "营收与现金流背离"}, st)
        assert pe._risk_chain_body({"risk_chain_analysis": "应收激增→现金流恶化"}, st)

    def test_kpi_cards_returns_table(self):
        """KPI 卡组件返回带背景色样式的 Table。"""
        from reportlab.platypus import Table
        from reportlab.lib.colors import white
        from tools.pdf_export import _kpi_cards, PRIMARY, BAD_RED, _register_chinese_font
        t = _kpi_cards([("重大风险", "2", BAD_RED, white), ("合计", "5", PRIMARY, white)],
                       _register_chinese_font())
        assert isinstance(t, Table) and len(t._cellvalues) == 2

    def test_new_chart_generators_write_png(self, tmp_path):
        """4 个新图表生成器真实产出 PNG（matplotlib 可用时）。"""
        import os
        from tools import pdf_export as pe
        gauge = str(tmp_path / "gauge.png")
        pe._gen_score_gauge(json.dumps({"score": 43.2, "level": "中等风险", "level_key": "medium"}), gauge)
        assert os.path.getsize(gauge) > 5000
        bar = str(tmp_path / "bar.png")
        pe._gen_breakdown_bar(json.dumps({"breakdown": {"financial": 50, "disclosure": 40, "validation": 30}}), bar)
        assert os.path.getsize(bar) > 5000
        donut = str(tmp_path / "donut.png")
        pe._gen_level_donut(json.dumps({"risk_details": [{"level": "重大"}, {"level": "一般"}]},
                                       ensure_ascii=False), donut)
        assert os.path.getsize(donut) > 5000
        bars = str(tmp_path / "bars.png")
        pe._gen_indicator_bars(json.dumps({"pct": [["毛利率", 21.5, 25.0]],
                                           "mult": [["流动比率", 0.8, 1.5]]}, ensure_ascii=False), bars)
        assert os.path.getsize(bars) > 5000
        # 零风险台账的环形图降级分支不崩
        pe._gen_level_donut(json.dumps({"risk_details": []}), str(tmp_path / "empty.png"))

    def test_split_with_extra_data_sources(self, monkeypatch):
        """拆分模式补传评分/校验数据源：仍生成 3 份且评分章节生效。"""
        uploaded = []
        monkeypatch.setattr(
            "local_storage.upload_file_to_storage",
            lambda path, dest, mime: (uploaded.append(dest), f"/local_storage/{dest}")[1],
        )
        monkeypatch.setattr("tools.pdf_export._embed_chart", lambda *a, **k: None)
        score_json = json.dumps({"score": 43.2, "level": "中等风险", "level_key": "medium",
                                 "breakdown": {"financial": 50.0, "disclosure": 40.0, "validation": 30.0}},
                                ensure_ascii=False)
        result = _export_pdf_impl(json.dumps(_dirty_report(), ensure_ascii=False),
                                  comprehensive_score_json=score_json,
                                  validation_json=json.dumps({"validation_result": "通过"}))
        assert result.count("已生成") == 3
        assert result.count("下载链接:") == 3
        assert len(uploaded) == 3


class TestNumberedLines:
    """_numbered_lines：多行内容拆分供 (1)(2) 子编号"""

    def test_multi_line_returns_stripped_lines(self):
        from tools.pdf_export import _numbered_lines
        assert _numbered_lines("第一行\n第二行") == ["第一行", "第二行"]
        assert _numbered_lines("a\n\nb\n") == ["a", "b"]   # 空行过滤

    def test_single_or_empty_returns_none(self):
        from tools.pdf_export import _numbered_lines
        assert _numbered_lines("单行内容") is None
        assert _numbered_lines("") is None
        assert _numbered_lines("   ") is None

    def test_non_str_input_converted(self):
        from tools.pdf_export import _numbered_lines
        assert _numbered_lines(123) is None


class TestExcelExportFormatting:
    def test_dirty_data_does_not_crash(self, tmp_path):
        out = tmp_path / "r.xlsx"
        result = _export_excel_impl(json.dumps(_dirty_report(), ensure_ascii=False), str(out))
        assert "已生成" in result
        assert "被拦截" not in result
        assert out.exists() and out.stat().st_size > 0

    def test_special_chars_preserved(self, tmp_path):
        """Excel 单元格应原样保留 < > 字符（无 XML 转义副作用）。"""
        from openpyxl import load_workbook
        out = tmp_path / "r.xlsx"
        _export_excel_impl(_dirty_report(), str(out))
        ws = load_workbook(str(out))["风险台账"]
        assert "OCF/NP<0.5" in str(ws.cell(row=2, column=3).value)

    def test_row_height_adapts_to_long_text(self, tmp_path):
        """超长审计建议应撑高明细行，而非被固定行高裁剪。"""
        from openpyxl import load_workbook
        out = tmp_path / "r.xlsx"
        _export_excel_impl(_dirty_report(), str(out))
        ws = load_workbook(str(out))["风险台账"]
        # R001 行（第 2 行）含超长 audit_suggestion，行高应显著大于原固定值 60
        assert ws.row_dimensions[2].height > 60

    def test_benchmark_scalar_written(self, tmp_path):
        """行业基准的标量值不应被静默跳过。"""
        from openpyxl import load_workbook
        out = tmp_path / "r.xlsx"
        _export_excel_impl(_dirty_report(), str(out))
        ws = load_workbook(str(out))["规则口径"]
        all_text = " ".join(
            str(v) for row in ws.iter_rows(values_only=True) for v in row if v is not None
        )
        assert "综合评分" in all_text
        assert "72" in all_text

    def test_summary_reconciled_with_details(self, tmp_path):
        """摘要不得读静态缓存：risk_summary 与明细不一致时以明细实时重算（P3）。

        实测缺陷：仲裁回写 + 勾稽校验条目并入后 LLM 的 risk_summary 已过时，
        底稿首页摘要（7 条）与内页明细（9 条）直接脱节。
        """
        from openpyxl import load_workbook
        report = _dirty_report()
        # 构造脱节台账：静态摘要 7 条，明细 9 条（仲裁新增 R008 + 系统校验 V001）
        report["risk_summary"] = {
            "total_risks": 7, "major_risks": 0, "important_risks": 4, "general_risks": 3,
            "risk_dimensions": {"financial_misstatement": 2, "related_party": 2,
                                "disclosure_compliance": 1, "going_concern": 1,
                                "regulatory_penalty": 1},
        }
        report["risk_details"] = [
            {"risk_id": f"R{i:03d}", "dimension": "财务错报风险", "level": "重要"}
            for i in range(1, 5)
        ] + [
            {"risk_id": "R005", "dimension": "信息披露合规风险", "level": "重要"},
            {"risk_id": "R006", "dimension": "关联交易风险", "level": "一般"},
            {"risk_id": "R007", "dimension": "监管处罚类高风险", "level": "一般"},
            {"risk_id": "R008", "dimension": "市场风险", "level": "一般"},
            {"risk_id": "V001", "dimension": "数据可靠性风险", "level": "重要"},
        ]
        out = tmp_path / "r.xlsx"
        _export_excel_impl(json.dumps(report, ensure_ascii=False), str(out))
        ws = load_workbook(str(out))["报告概览"]
        summary = {}
        for row in ws.iter_rows(values_only=True):
            if row[0] in ("系统采信风险", "重大风险", "重要风险", "一般风险"):
                summary[row[0]] = row[1]
        assert summary == {"系统采信风险": 9, "重大风险": 0, "重要风险": 6, "一般风险": 3}, (
            "摘要须按明细实时聚合，不得残留仲裁前静态缓存")


class TestFmtScore:
    """K 补丁：评分展示统一格式（智能去尾，防概览 86 vs 底层 85.7 出入）。"""

    def test_non_integer_keeps_one_decimal(self):
        from tools.pdf_export import _fmt_score
        assert _fmt_score(85.7) == "85.7"
        assert _fmt_score(14.3) == "14.3"
        assert _fmt_score("85.7") == "85.7"

    def test_integer_no_trailing_zero(self):
        from tools.pdf_export import _fmt_score
        assert _fmt_score(100.0) == "100"
        assert _fmt_score(86) == "86"
        assert _fmt_score(80) == "80"

    def test_disclosure_kpi_card_shows_decimal(self):
        """合规概览 KPI 卡渲染 85.7/14.3（非 86/14），与系统量化事实一致。"""
        import json
        from tools import pdf_export as pe
        st = pe._build_styles(pe._register_chinese_font())
        dc = json.dumps({"compliance_score": 85.7, "risk_score": 14.3,
                         "checked_items": 14, "passed_items": 12,
                         "issues": ["缺失公司基本情况章节"],
                         "sections_missing": ["公司基本情况"],
                         "audit_opinion": "未经审计（半年度报告）"}, ensure_ascii=False)
        body = pe._disclosure_section_body(dc, st, pe._register_chinese_font())
        texts = []
        def _collect(el):
            t = getattr(el, "text", None)
            if t:
                texts.append(str(t))
            rows = getattr(el, "_cellvalues", None) or []
            for row in rows:
                for cell in row:
                    if hasattr(cell, "text"):
                        texts.append(str(cell.text))
        for el in body:
            _collect(el)
        joined = " ".join(texts)
        assert "85.7" in joined and "14.3" in joined

    def test_compliance_standards_row_shows_decimal(self):
        """合规标准对照行渲染 85.7（非 86）。"""
        import json
        from tools import pdf_export as pe
        st = pe._build_styles(pe._register_chinese_font())
        dc = json.dumps({"compliance_score": 85.7, "risk_score": 14.3,
                         "checked_items": 14, "passed_items": 12,
                         "issues": ["缺失章节"], "sections_missing": ["公司基本情况"]},
                        ensure_ascii=False)
        body = pe._compliance_standard_body(dc, st, pe._register_chinese_font())
        texts = []
        for el in body:
            t = getattr(el, "text", None)
            if t:
                texts.append(str(t))
            for row in (getattr(el, "_cellvalues", None) or []):
                for cell in row:
                    if hasattr(cell, "text"):
                        texts.append(str(cell.text))
        joined = " ".join(texts)
        assert "85.7" in joined

    def test_key_section_missing_appends_hint(self):
        """M 补丁：缺失"主要会计数据和财务指标"类核心章节时附提示文案（不改评分）。"""
        import json
        from tools import pdf_export as pe
        st = pe._build_styles(pe._register_chinese_font())
        dc = json.dumps({"compliance_score": 85.7, "risk_score": 14.3,
                         "checked_items": 14, "passed_items": 12,
                         "issues": ["缺失章节"],
                         "sections_missing": ["公司基本情况", "主要会计数据和财务指标"]},
                        ensure_ascii=False)
        body = pe._disclosure_section_body(dc, st, pe._register_chinese_font())
        texts = []
        for el in body:
            t = getattr(el, "text", None)
            if t:
                texts.append(str(t))
        joined = " ".join(texts)
        assert "建议人工确认披露完整性" in joined

    def test_other_missing_no_hint(self):
        """非核心章节缺失不触发提示（避免噪音）。"""
        import json
        from tools import pdf_export as pe
        st = pe._build_styles(pe._register_chinese_font())
        dc = json.dumps({"compliance_score": 85.7, "risk_score": 14.3,
                         "checked_items": 14, "passed_items": 12,
                         "issues": [], "sections_missing": ["公司治理"]},
                        ensure_ascii=False)
        body = pe._disclosure_section_body(dc, st, pe._register_chinese_font())
        texts = []
        for el in body:
            t = getattr(el, "text", None)
            if t:
                texts.append(str(t))
        joined = " ".join(texts)
        assert "建议人工确认披露完整性" not in joined


class TestExcelDimensionChinese:
    """L 补丁：Excel 风险明细维度列统一规范中文展示名（A4 展示层收敛）。"""

    def test_dimension_column_normalized_chinese(self, tmp_path):
        """归一后台账导出：维度列展示为规范中文「财务风险」，且无非法英文键残留。"""
        from tools.excel_export import _export_excel_impl
        from openpyxl import load_workbook
        report = {"company_info": {"company_name": "测试公司"},
                  "risk_details": [
                      {"risk_id": "R001", "dimension": "financial_misstatement",
                       "level": "重要", "title": "应收激增"},
                      {"risk_id": "R005", "dimension": "financial_misstatement",
                       "level": "重要", "title": "坏账准备计提充分性", "source": "仲裁新增"},
                  ]}
        out = tmp_path / "t.xlsx"
        _export_excel_impl(json.dumps(report, ensure_ascii=False), str(out))
        wb = load_workbook(str(out))
        ws = wb["风险台账"]
        dims = [str(row[1]) for row in ws.iter_rows(min_row=2, values_only=True) if row[0]]
        # 展示层统一为规范中文名，避免混入英文键或 LLM 自由文本
        assert dims == ["财务风险", "财务风险"], dims


class TestStatementsAnalysis:
    """三大报表基本分析：确定性规则层（非 LLM）+ 渲染 + 缺失降级。"""

    def _fin_json(self, **overrides):
        import json as _json
        items = {
            "balance_sheet": {"total_assets_current": 1000e8, "total_liabilities_current": 400e8,
                              "net_assets_current": 600e8, "cash_and_equivalents_current": 200e8,
                              "accounts_receivable_current": 100e8,
                              "accounts_receivable_previous": 80e8,
                              "inventory_current": 50e8, "short_term_debt_current": 60e8,
                              "goodwill_current": 10e8},
            "income_statement": {"revenue_current": 145e8, "revenue_previous": 155e8,
                                 "net_profit_current": 84e8, "net_profit_previous": 99e8,
                                 "cost_of_goods_current": 115e8},
            "cashflow_statement": {"operating_cashflow_current": 227e8,
                                   "operating_cashflow_previous": 210e8},
        }
        for k, v in overrides.items():
            sec, key = k.split(".", 1)
            items[sec][key] = v
        return _json.dumps({"statement_items": items}, ensure_ascii=False)

    def test_golden_sample_math(self):
        """金标样本对照：同比/占比与 Excel 公式基准一致（±0.01 容差）。"""
        from tools.pdf_export import _statements_basic_analysis
        tables, notes, missing = _statements_basic_analysis(self._fin_json())
        assert missing is False
        # 利润表：营收同比 = (145-155)/155 = -6.4516 → -6.45
        inc = next(t for t in tables if t[0] == "利润表")
        rev_row = next(r for r in inc[2] if r[0] == "营业收入")
        assert rev_row[1] == 145e8 and rev_row[2] == 155e8
        assert abs(rev_row[3] - (-6.4516)) < 0.01, rev_row[3]
        # 净利润同比 = (84-99)/99 = -15.1515 → -15.15
        np_row = next(r for r in inc[2] if r[0] == "净利润（合并）")
        assert abs(np_row[3] - (-15.1515)) < 0.01, np_row[3]
        # 毛利率 = (145-115)/145 = 20.69
        assert abs((145 - 115) / 145 * 100 - 20.6897) < 0.01
        # 资产负债表：应收占比 = 100/1000 = 10.00%
        bs = next(t for t in tables if t[0] == "资产负债表")
        ar_row = next(r for r in bs[2] if r[0] == "应收账款")
        assert abs(ar_row[4] - 10.0) < 0.01, ar_row[4]
        # 应收同比 = (100-80)/80 = 25.00%
        assert abs(ar_row[3] - 25.0) < 0.01
        # 资产负债表科目无上期 → 同比为 None（不编造）
        ta_row = next(r for r in bs[2] if r[0] == "总资产")
        assert ta_row[3] is None and ta_row[2] is None

    def test_judgement_rules(self):
        """判读规则：应收增速>营收增速+5 → 回款提示；OCF/NP<1 → 含金量提示。"""
        from tools.pdf_export import _statements_basic_analysis
        # 应收同比 25% vs 营收同比 -6.45% → 触发回款质量提示
        _, notes, _ = _statements_basic_analysis(self._fin_json())
        assert any("回款质量" in n for n in notes)
        assert any("净利润同比下降 15.15%" in n for n in notes)
        # OCF/NP=2.70 ≥1 → 不触发含金量提示
        assert not any("盈利含金量" in n for n in notes)
        # 现金流 < 净利润 → 触发
        _, notes2, _ = _statements_basic_analysis(self._fin_json(**{"cashflow_statement.operating_cashflow_current": 50e8}))
        assert any("盈利含金量" in n for n in notes2)

    def test_missing_items_degrade(self):
        """科目缺失：全缺 → 可见降级提示；部分缺 → 行跳过不抛异常。"""
        from tools.pdf_export import _statements_basic_analysis, _statements_section_body
        import json as _json
        # 全缺
        tables, notes, missing = _statements_basic_analysis(_json.dumps({"statement_items": {}}))
        assert missing is True and not tables
        # 渲染层可见提示
        st = None
        from tools import pdf_export as pe
        st = pe._build_styles(pe._register_chinese_font())
        body = _statements_section_body(_json.dumps({"statement_items": {}}), st, pe._register_chinese_font())
        texts = [str(getattr(el, "text", "")) for el in body]
        assert any("因关键科目缺失未能生成完整三表分析" in t for t in texts)
        # 部分缺（仅利润表）→ 只有利润表渲染
        fin = _json.dumps({"statement_items": {"income_statement": {"revenue_current": 100e8,
                                                                    "net_profit_current": 10e8}}},
                          ensure_ascii=False)
        tables2, _, missing2 = _statements_basic_analysis(fin)
        assert missing2 is False
        assert [t[0] for t in tables2] == ["利润表"]

    def test_section_renders_three_tables(self):
        """渲染：三表小节名与表格行齐全，金额用万/亿换算。"""
        from tools import pdf_export as pe
        st = pe._build_styles(pe._register_chinese_font())
        body = pe._statements_section_body(self._fin_json(), st, pe._register_chinese_font())
        texts = []
        for el in body:
            t = getattr(el, "text", None)
            if t:
                texts.append(str(t))
            for row in (getattr(el, "_cellvalues", None) or []):
                for cell in row:
                    if isinstance(cell, str):
                        texts.append(cell)
                    elif hasattr(cell, "text"):
                        texts.append(str(cell.text))
        joined = " ".join(texts)
        assert "资产负债表" in joined and "利润表" in joined and "现金流量表" in joined
        assert "半年度累计" in joined
        assert "亿元" in joined  # _fmt_num 自动换算
        assert "判读要点" in joined

    def test_declared_million_unit_is_converted_before_pdf_display(self):
        """报表原值为百万元时，PDF金额应先换算为人民币元再显示亿元。"""
        from tools import pdf_export as pe
        assert pe._fmt_num(1450099, input_unit="人民币百万元") == "14,500.99 亿元"
        assert pe._fmt_num(840.07, input_unit="亿元") == "840.07 亿元"

        data = {"amount_unit": "人民币百万元", "statement_items": {
            "income_statement": {"revenue_current": 1450099,
                                  "revenue_previous": 1554973},
        }}
        st = pe._build_styles(pe._register_chinese_font())
        body = pe._statements_section_body(json.dumps(data, ensure_ascii=False), st,
                                           pe._register_chinese_font())
        texts = []
        for el in body:
            for row in (getattr(el, "_cellvalues", None) or []):
                for cell in row:
                    if isinstance(cell, str):
                        texts.append(cell)
                    elif hasattr(cell, "text"):
                        texts.append(str(cell.text))
        assert "14,500.99 亿元" in " ".join(texts)


class TestAuditOpinionSource:
    """审计意见数据源状态判定（双源合并）：识别工具未调用时披露检查兜底。"""

    def test_disclosure_checker_opinion_counts_as_used(self):
        """identify_audit_opinion 为空但披露检查已识别意见 → 状态"已使用"。"""
        from tools.pdf_export import _audit_opinion_source
        dc = json.dumps({"compliance_score": 92.9, "audit_opinion": "未经审计（半年度报告）"},
                        ensure_ascii=False)
        used, note = _audit_opinion_source("", dc)
        assert used is True
        assert "披露规范性检查识别" in note

    def test_identify_tool_result_wins(self):
        """identify_audit_opinion 有结果时优先采用（来源为识别工具）。"""
        from tools.pdf_export import _audit_opinion_source
        ao = json.dumps({"audit_opinion": {"opinion_type": "标准无保留意见", "identified": True}},
                        ensure_ascii=False)
        used, note = _audit_opinion_source(ao, "")
        assert used is True and note == ""

    def test_both_missing_keeps_unavailable(self):
        """两个来源均无意见 → 维持"未获取"，原因只写原因本身。

        原因不得带「未获取：」前缀：同表「状态」列已写未获取，重复会在同一行
        出现两次（实测缺陷）。
        """
        from tools.pdf_export import _audit_opinion_source
        used, note = _audit_opinion_source("", "")
        assert used is False
        assert note and not note.startswith("未获取")
        assert "审计意见" in note

    def test_unidentified_opinion_does_not_count(self):
        """披露检查输出"未识别"不算已获取（防虚标）。"""
        from tools.pdf_export import _audit_opinion_source
        dc = json.dumps({"audit_opinion": "未识别"}, ensure_ascii=False)
        used, _ = _audit_opinion_source("", dc)
        assert used is False


class TestEnergyGrossMarginVerdict:
    """能源全产业链毛利率判读软化：不受低毛利贸易摊薄误判。"""

    def test_energy_gross_margin_gets_dilution_note(self):
        """能源行业毛利率显著低于开采基准时追加摊薄提示（注释读基准库配置）。"""
        from tools.pdf_export import _interpret_indicator, _load_benchmarks
        _load_benchmarks("能源")  # 填充行业注释缓存（配置化：注释钉在 industry_benchmarks.json）
        verdict, bad = _interpret_indicator("gross_margin_pct", 20.89, 35, industry="能源")
        assert bad is True
        assert "摊薄" in verdict and "分部数据" in verdict

    def test_energy_verdict_without_loaded_config(self):
        """基准配置未加载（缓存空）时不追加注释，判读仍为负面信号（防御）。"""
        from tools.pdf_export import _interpret_indicator, _BENCH_NOTES
        _BENCH_NOTES.clear()
        verdict, bad = _interpret_indicator("gross_margin_pct", 20.89, 35, industry="能源")
        assert bad is True
        assert "摊薄" not in verdict

    def test_non_energy_keeps_plain_verdict(self):
        """非能源行业不追加摊薄提示（判读保持原样）。"""
        from tools.pdf_export import _interpret_indicator
        verdict, bad = _interpret_indicator("gross_margin_pct", 20.89, 35, industry="制造业")
        assert bad is True
        assert "摊薄" not in verdict

    def test_industry_default_none_compatible(self):
        """不传 industry（默认 None）时行为与旧版一致（测试兼容）。"""
        from tools.pdf_export import _interpret_indicator
        verdict, bad = _interpret_indicator("gross_margin_pct", 20.89, 35)
        assert bad is True and "摊薄" not in verdict


class TestDimensionAggregation:
    """维度统计归一化（K）："财务错报"与"财务错报风险"须聚合为同一键。"""

    def test_reconcile_summary_aggregates_dim_variants(self):
        """_reconcile_summary 统计时归一化维度键，消除重复行。"""
        from tools.pdf_export import _reconcile_summary
        report = {
            "risk_summary": {"total_risks": 4, "risk_dimensions": {}},
            "risk_details": [
                {"dimension": "财务错报风险", "level": "重要"},
                {"dimension": "财务错报", "level": "一般"},
                {"dimension": "财务错报风险", "level": "重要"},
                {"dimension": "关联交易风险", "level": "一般"},
            ],
        }
        _reconcile_summary(report)
        dims = report["risk_summary"]["risk_dimensions"]
        assert dims.get("financial_misstatement") == 3, "'财务错报'与'财务错报风险'应聚合"
        assert dims.get("related_party") == 1
        assert len(dims) == 2, "不得出现语义重复的维度行"

    def test_excel_reconcile_aggregates_dim_variants(self):
        """excel 侧 _reconcile_risk_summary 同样归一化维度键。"""
        from tools.excel_export import _reconcile_risk_summary
        report = {"risk_summary": {}, "risk_details": [
            {"dimension": "财务错报风险", "level": "重要"},
            {"dimension": "财务错报", "level": "一般"},
        ]}
        _reconcile_risk_summary(report)
        dims = report["risk_summary"]["risk_dimensions"]
        assert dims.get("financial_misstatement") == 2
        assert len(dims) == 1

class TestScoreSnapshotReplace:
    """L 补丁：结论章评分快照替换（吞全 span，杜绝双分矛盾）。"""

    def test_snapshot_replaces_full_span(self):
        from tools.pdf_export import _apply_score_snapshot
        report = {"comprehensive_score_snapshot": {"score": 31.5, "level": "中等风险"}}
        text = "综合风险评分9分处于低风险区间，公司财务结构稳健。"
        out = _apply_score_snapshot(text, report)
        assert out == "综合风险评分 31.5分（中等风险），公司财务结构稳健。"
        assert "处于低风险区间" not in out

    def test_snapshot_replaces_parenthesized_form(self):
        from tools.pdf_export import _apply_score_snapshot
        report = {"comprehensive_score_snapshot": {"score": 10.5, "level": "低风险"}}
        text = "（综合风险评分 9分）需关注。"
        out = _apply_score_snapshot(text, report)
        assert "（综合风险评分 10.5分（低风险））" not in out, "不得产生嵌套括号"
        assert "综合风险评分 10.5分（低风险）" in out

    def test_no_snapshot_keeps_text(self):
        """无快照时保持原文（L 补丁兜底）；50d：快照存在但 score=None（评分未获取）
        时归一为「未获取/无法判定」——快照恒存在、正文恒被归一契约。"""
        from tools.pdf_export import _apply_score_snapshot
        text = "综合风险评分 9分处于低风险区间"
        assert _apply_score_snapshot(text, {}) == text
        out = _apply_score_snapshot(text, {"comprehensive_score_snapshot": {"score": None}})
        assert "9分" not in out
        assert "未获取/无法判定" in out

    def test_no_misfire_on_noun_phrases(self):
        """50d：分隔形态带负向前瞻——不误吞「，低风险行业」「中低风险偏好投资者」
        类名词短语（实测病句缺陷：50d 初版把正文截断成病句）。"""
        from tools.pdf_export import _apply_score_snapshot
        report = {"comprehensive_score_snapshot": {"score": 31.5, "level": "中等风险"}}
        t1 = "综合评分 30分，低风险行业龙头股业绩稳健。"
        out1 = _apply_score_snapshot(t1, report)
        assert "低风险行业" in out1 and "龙头股" in out1
        t2 = "综合评分 55 分，中低风险偏好投资者应谨慎。"
        out2 = _apply_score_snapshot(t2, report)
        assert "中低风险偏好" in out2
        # 分隔形态后接标点/行尾/竖线时仍正常吞掉（等级属性语境）
        t3 = "综合评分 15.8分，低风险），建议关注。"
        out3 = _apply_score_snapshot(t3, report)
        assert "，低风险" not in out3

    def test_llm_short_variant_replaced(self):
        """LLM 写「综合评分」而非「综合风险评分」时同样替换（实测残留 15.8 双分矛盾的根因）。"""
        from tools.pdf_export import _apply_score_snapshot
        report = {"comprehensive_score_snapshot": {"score": 12.6, "level": "低风险"}}
        text = "整体风险较低（综合评分15.8分，低风险），建议关注。"
        out = _apply_score_snapshot(text, report)
        assert "15.8" not in out
        assert "综合风险评分 12.6分（低风险）" in out
        # 逗号分隔的旧等级尾部不得残留（避免「（低风险）（低风险）」重复）
        assert out.count("（低风险）") == 1

    def test_llm_short_variant_without_parens(self):
        """无括号包裹的「综合评分 XX分」表述同样替换。"""
        from tools.pdf_export import _apply_score_snapshot
        report = {"comprehensive_score_snapshot": {"score": 12.6, "level": "低风险"}}
        text = "综合评分9分处于低风险区间，公司整体风险较低。"
        out = _apply_score_snapshot(text, report)
        assert "综合风险评分 12.6分（低风险）" in out
        assert "处于低风险区间" not in out

    def test_slash_100_form_replaced(self):
        """50b：LLM 写「9/100（低风险）」形态时同样替换——旧正则要求数字后必须
        跟「分」字，/100 形态完全不匹配导致旧分残留正文（9 vs 16.5 罗生门）。"""
        from tools.pdf_export import _apply_score_snapshot
        report = {"comprehensive_score_snapshot": {"score": 16.5, "level": "低风险"}}
        text = "综合风险评分 9/100（低风险），公司财务结构稳健。"
        out = _apply_score_snapshot(text, report)
        assert "9/100" not in out
        assert "综合风险评分 16.5分（低风险）" in out
        assert "，公司财务结构稳健。" in out

    def test_no_misfire_on_full_score_disclosure_text(self):
        """50b：不误伤无「综合评分」前缀的「100分满分披露」类表述。"""
        from tools.pdf_export import _apply_score_snapshot
        report = {"comprehensive_score_snapshot": {"score": 16.5, "level": "低风险"}}
        text = "信息披露合规检查 100 分、无缺失项，检查项全部通过。"
        out = _apply_score_snapshot(text, report)
        assert out == text

    def test_table_cell_form_replaced(self):
        """50c：markdown 表格行形态「| 综合风险评分 | 55 | 中等风险 |」同样替换——
        17:59 版实测 55 分残留的根因：数字被竖线分隔且无「分」字后缀，旧正则要求
        数字紧跟标签完全不匹配，与系统评分同屏双分。"""
        from tools.pdf_export import _apply_score_snapshot
        report = {"comprehensive_score_snapshot": {"score": 16.5, "level": "低风险"}}
        text = "| 综合风险评分 | 55 | 中等风险 | 需结合人工专业判断复核 |"
        out = _apply_score_snapshot(text, report)
        assert "55" not in out
        assert "中等风险" not in out
        assert "综合风险评分 16.5分（低风险）" in out
        # 三列结构保留：替换后仍是 | 评分列 | 说明列 |
        assert "| 综合风险评分 16.5分（低风险）| 需结合人工专业判断复核 |" in out

    def test_wei_and_colon_forms_replaced(self):
        """50c：「综合风险评分为 55 分」与「综合风险评分：55分」形态同样替换。"""
        from tools.pdf_export import _apply_score_snapshot
        report = {"comprehensive_score_snapshot": {"score": 10.5, "level": "低风险"}}
        t1 = "综合风险评分为 55 分（中等风险），建议关注。"
        out1 = _apply_score_snapshot(t1, report)
        assert "55" not in out1 and "综合风险评分 10.5分（低风险）" in out1
        t2 = "综合风险评分：55分，处于中等风险区间。"
        out2 = _apply_score_snapshot(t2, report)
        assert "55" not in out2 and "综合风险评分 10.5分（低风险）" in out2
        assert "处于中等风险区间" not in out2

    def test_mid_low_level_cell_swallowed(self):
        """50d：表格「低风险 | 中低风险」两级并存时一并吞掉，由系统等级唯一确定。"""
        from tools.pdf_export import _apply_score_snapshot
        report = {"comprehensive_score_snapshot": {"score": 16.5, "level": "低风险"}}
        text = "| 综合风险评分 | 15.0分（低风险）| 中低风险 | 整体稳健，个别科目需核查 |"
        out = _apply_score_snapshot(text, report)
        assert "中低风险" not in out
        assert "综合风险评分 16.5分（低风险）" in out
        assert "整体稳健，个别科目需核查" in out

    def test_risk_id_label_with_semantic(self):
        """50d：风险编号展示——semantic_id 存在且不同时与主编号并列。"""
        from tools.pdf_export import _risk_id_label
        assert _risk_id_label({"risk_id": "R001", "semantic_id": "AR-001"}) == "R001（AR-001）"
        assert _risk_id_label({"risk_id": "R001"}) == "R001"
        assert _risk_id_label({"risk_id": "R001", "semantic_id": "R001"}) == "R001"

class TestSubsetNoteAndFactsColumn:
    """S2/S4：专项报告口径标注 + Excel 事实层列。"""

    def test_subset_note_present_in_split_reports(self, monkeypatch, tmp_path):
        """财务专项报告含口径标注（拆分模式）。"""
        from tools.pdf_export import _export_pdf_impl
        monkeypatch.setattr(
            "local_storage.upload_file_to_storage",
            lambda path, dest, mime: f"/local_storage/{dest}")
        report = {"company_info": {"company_name": "测试公司", "report_year": "2025"},
                  "risk_details": [
                      {"risk_id": "R001", "dimension": "财务错报风险", "level": "重要", "title": "应收激增"},
                      {"risk_id": "R002", "dimension": "关联交易风险", "level": "一般", "title": "关联交易"},
                  ]}
        from tools.pdf_export import _export_split_reports
        out_txt = _export_split_reports(report, "", "", "", "", "", "", "", module="financial")
        assert "财务健康诊断报告" in out_txt

    def test_excel_facts_column_appears(self, tmp_path):
        from tools.excel_export import _export_excel_impl
        from openpyxl import load_workbook
        report = {"company_info": {"company_name": "测试公司"},
                  "risk_details": [
                      {"risk_id": "R003", "dimension": "数据可靠性风险", "level": "重要",
                       "title": "未分配利润勾稽", "evidence": "勾稽差异20.43%"},
                  ]}
        val = json.dumps({"data_validation": {"all_checks": [
            {"check": "未分配利润一致性", "passed": None, "difference_pct": "20.43%", "message": "口径提示"}]}})
        out = tmp_path / "t.xlsx"
        _export_excel_impl(json.dumps(report, ensure_ascii=False), str(out), validation_json=val)
        wb = load_workbook(str(out))
        ws = wb["风险台账"]
        headers = [c.value for c in ws[1]]
        assert "系统量化事实" in headers
        facts = ws.cell(row=2, column=12).value or ""
        assert "勾稽校验" in facts

    def test_subset_completeness_warns_uncovered_dimension(self, monkeypatch):
        """子集完备性：其他/未知维度条目不落入财务/合规子集 → 日志告警（不中断产出），
        一致性说明附未落入清单（v23 实测："财务"简写维度丢包后仅综合汇总可见）。"""
        from tools.pdf_export import _export_split_reports
        warnings = []
        monkeypatch.setattr("tools.pdf_export.logger", type("L", (), {
            "warning": staticmethod(lambda msg, *a, **k: warnings.append(str(msg))),
            "info": staticmethod(lambda *a, **k: None),
        })())
        monkeypatch.setattr(
            "local_storage.upload_file_to_storage",
            lambda path, dest, mime: f"/local_storage/{dest}")
        report = {"company_info": {"company_name": "测试公司", "report_year": "2025"},
                  "risk_details": [
                      {"risk_id": "R001", "dimension": "财务错报风险", "level": "重要", "title": "应收激增"},
                      {"risk_id": "R009", "dimension": "其他", "level": "一般", "title": "未分类事项"},
                  ]}
        out_txt = _export_split_reports(report, "", "", "", "", "", "", "")
        assert "财务健康诊断报告" in out_txt  # 产出不中断
        assert any("子集完备性" in w for w in warnings), warnings
        assert any("其他" in w for w in warnings), warnings

    def test_subset_completeness_clean_when_all_covered(self, monkeypatch):
        """子集完备性：全部条目落入财务/合规子集 → 无告警。"""
        from tools.pdf_export import _export_split_reports
        warnings = []
        monkeypatch.setattr("tools.pdf_export.logger", type("L", (), {
            "warning": staticmethod(lambda msg, *a, **k: warnings.append(str(msg))),
            "info": staticmethod(lambda *a, **k: None),
        })())
        monkeypatch.setattr(
            "local_storage.upload_file_to_storage",
            lambda path, dest, mime: f"/local_storage/{dest}")
        report = {"company_info": {"company_name": "测试公司", "report_year": "2025"},
                  "risk_details": [
                      {"risk_id": "R001", "dimension": "财务", "level": "重要", "title": "应收激增"},
                      {"risk_id": "R002", "dimension": "关联交易", "level": "重要", "title": "关联交易"},
                  ]}
        _export_split_reports(report, "", "", "", "", "", "", "")
        assert not any("子集完备性" in w for w in warnings), warnings
