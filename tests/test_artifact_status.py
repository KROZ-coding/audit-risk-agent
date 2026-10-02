"""产物三态与任务级「部分完成」回归测试（含网页分区与图表等级色）。

整改要求：每个 PDF、Excel、图表分别记录生成中、成功或失败，任务可显示部分完成；
实际文件存在且可访问才提供下载。网页分区顺序为「报告元数据与完整性 → 四项能力及
补充分析 → 正式风险 → 待补证据与待复核 → 产物下载」。图表等级色为重大红、重要黄、
一般蓝，且保留文字标识。
"""
import json
import os
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

from core.result_contract import (
    ARTIFACT_FAILED,
    ARTIFACT_GENERATING,
    ARTIFACT_SUCCESS,
    TASK_COMPLETED,
    TASK_FAILED,
    TASK_NOT_RUN,
    TASK_PARTIAL,
    TASK_RUNNING,
    ArtifactManifest,
    derive_task_status,
)
from main import expected_artifacts

_WEB = os.path.join(os.path.dirname(__file__), "..", "src", "web", "index.html")


def _artifact(status: str) -> dict:
    return ArtifactManifest(artifact_id="a", kind="pdf", status=status).to_dict()


class TestTaskStatus:
    """任务级状态：必须能表达「部分完成」，且与三态清单一致。"""

    def test_generating_means_running(self):
        assert derive_task_status([_artifact(ARTIFACT_SUCCESS),
                                   _artifact(ARTIFACT_GENERATING)]) == TASK_RUNNING

    def test_all_success_completed(self):
        assert derive_task_status([_artifact(ARTIFACT_SUCCESS),
                                   _artifact(ARTIFACT_SUCCESS)]) == TASK_COMPLETED

    def test_mixed_results_partial(self):
        assert derive_task_status([_artifact(ARTIFACT_SUCCESS),
                                   _artifact(ARTIFACT_FAILED)]) == TASK_PARTIAL

    def test_all_failed(self):
        assert derive_task_status([_artifact(ARTIFACT_FAILED)]) == TASK_FAILED

    def test_empty_manifest_not_run(self):
        assert derive_task_status([]) == TASK_NOT_RUN

    def test_manifest_keeps_three_states(self):
        assert ArtifactManifest(artifact_id="a", kind="chart",
                                status=ARTIFACT_GENERATING).to_dict()["status"] == "generating"


class TestExpectedArtifacts:
    """运行开始时登记的期望清单（用于「生成中」展示，与 Agent 侧同构）。"""

    def test_full_pipeline_lists_all(self):
        keys = [item["key"] for item in expected_artifacts(None, False)]
        assert keys == ["heatmap", "radar", "trend", "pdf_financial", "pdf_compliance",
                        "pdf_synthesis", "excel", "json"]

    def test_fast_mode_skips_charts(self):
        keys = [item["key"] for item in expected_artifacts("financial", True)]
        assert keys == ["pdf_financial", "excel", "json"]

    def test_compliance_module_only_compliance_pdf(self):
        keys = [item["key"] for item in expected_artifacts("compliance", False)]
        assert "pdf_compliance" in keys and "pdf_financial" not in keys

    def test_outlook_has_no_file_artifacts(self):
        assert expected_artifacts("outlook", False) == []


class TestManifestPayload:
    """最终报告：清单逐项三态 + 任务级状态 + 版本可追溯。"""

    def _build(self):
        from main import GraphService

        service = object.__new__(GraphService)
        messages = [
            {"type": "tool", "name": "generate_risk_heatmap",
             "content": '{"download_url": "/local_storage/charts/h.png"}'},
            {"type": "tool", "name": "export_excel_report",
             "content": "Excel已生成: /local_storage/files/e.xlsx"},
            {"type": "ai", "content": json.dumps({
                "company_info": {"company_name": "产物测试公司"},
                "risk_details": []}, ensure_ascii=False)},
        ]
        return service._build_final_report_from_messages(messages)

    def test_partial_task_status_and_downloads(self, monkeypatch, tmp_path):
        import main as main_mod
        # 热力图与 Excel 真实存在，三份 PDF 不存在 → 部分完成
        heat = tmp_path / "h.png"
        heat.write_bytes(b"png")
        excel = tmp_path / "e.xlsx"
        excel.write_bytes(b"xlsx")
        monkeypatch.setattr(main_mod, "_artifact_is_accessible", lambda p: True)
        monkeypatch.setattr(main_mod, "_artifact_key", lambda p: {
            "h.png": "heatmap", "e.xlsx": "excel"}.get(os.path.basename(str(p)), ""))
        report = self._build()
        assert report["task_status"] == TASK_PARTIAL
        items = {item["name"]: item for item in report["artifact_manifest"]}
        assert items["h.png"]["accessible"] is True
        assert items["h.png"]["url"]
        failed = [item for item in report["artifact_manifest"] if item["status"] == "failed"]
        assert failed and all(not item["accessible"] for item in failed)
        assert all(item["data_version"] and item["rule_version"]
                   for item in report["artifact_manifest"])

    def test_report_metadata_and_sources(self):
        report = self._build()
        meta = report["report_metadata"]
        assert meta["company_name"] == "产物测试公司"
        assert meta["data_version"] and meta["rule_version"]
        assert meta["source_total"] == 7
        missing = [item for item in meta["data_sources"] if not item["used"]]
        assert missing and all(item["note"] for item in missing)
        assert report["indicator_view"]["available"] is False


class TestWebSections:
    """网页分区与产物三态渲染入口存在，且新事件有处理分支。"""

    @pytest.fixture(scope="class")
    def html(self):
        with open(_WEB, encoding="utf-8") as fh:
            return fh.read()

    def test_metadata_section_present(self, html):
        assert "报告元数据与完整性" in html
        assert "report.report_metadata" in html
        assert "数据来源与完整性说明" in html

    def test_capability_section_present(self, html):
        assert "财务指标与审计关注分析（指标计算过程与依据）" in html
        assert "report.indicator_view" in html
        assert "代入过程（含结果）" in html

    def test_artifact_three_states_and_partial(self, html):
        assert "generating: '生成中'" in html
        assert "partial: '部分完成'" in html

    def test_artifact_status_event_handled(self, html):
        assert html.count("data.type === 'artifact_status'") == 2
        assert "function updateArtifactProgress" in html

    def test_review_budget_visible(self, html):
        assert "gate.review_budget" in html
        assert "转人工复核" in html


class TestLevelColors:
    """图表等级色：重大红、重要黄、一般蓝，并保留文字标识。"""

    @staticmethod
    def _hex(color) -> str:
        return "#" + str(color.hexval()).lstrip("#").upper().removeprefix("0X")

    def test_pdf_level_colors_are_planned_colors(self):
        from tools import pdf_export as pe
        assert self._hex(pe.LEVEL_COLORS["重大"]) == "#FF0000"
        assert self._hex(pe.LEVEL_COLORS["重要"]) == "#FFD400"
        assert self._hex(pe.LEVEL_COLORS["一般"]) == "#4472C4"
        # 白底文字用深黄（同色系），不再退回橙色
        assert self._hex(pe.LEVEL_TEXT_COLORS["重要"]) == "#B8860B"
        assert self._hex(pe.LEVEL_COLORS["重要"]) != "#FF8C00"

    def test_excel_level_fill_yellow_with_dark_font(self, tmp_path):
        from openpyxl import load_workbook

        from tools.excel_export import _export_excel_impl

        report = {"company_info": {"company_name": "等级色测试"},
                  "risk_details": [{"risk_id": "R001", "dimension": "财务错报风险",
                                    "title": "测试", "level": "重要"}]}
        out = tmp_path / "r.xlsx"
        _export_excel_impl(json.dumps(report, ensure_ascii=False), str(out))
        ws = load_workbook(str(out))["风险台账"]
        cell = ws.cell(row=2, column=4)
        assert cell.value == "重要"
        assert (cell.fill.start_color.rgb or "").endswith("FFD400")
        assert (cell.font.color.rgb or "").endswith("3F3000")

    def test_heatmap_uses_level_columns(self):
        """热力图按等级列着色（不再用连续色阶），并保留格内数量文字。"""
        import inspect

        from tools import visualizer
        source = inspect.getsource(visualizer._generate_risk_heatmap)
        assert "YlOrRd" not in source
        assert "level_rgb" in source
        for color in ("'重大': (1.0, 0.0, 0.0)", "'重要': (1.0, 0.83, 0.0)",
                      "'一般': (0.27, 0.45, 0.77)"):
            assert color in source
