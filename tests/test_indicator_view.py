"""指标计算过程（代入过程 / 原文定位 / 四项能力与补充分析）回归测试。

整改要求：每个用于判断的指标展示「公式→输入值及单位→期间与口径→代入过程→
结果→阈值或基准及来源→状态→原文定位」。代入过程必须由确定性计算生成、可复算，
缺失输入与零分母如实标注，不得以中性值补齐，也不得由模型转述。
"""
import json
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

from tools.financial_calculator import calculate_financial_indicators
from tools.indicator_view import build_indicator_view


def _data(**overrides) -> dict:
    data = {
        "amount_unit": "万元", "period": "2025年度", "scope": "合并",
        "revenue_current": 120000, "revenue_previous": 100000,
        "net_profit_current": 8000, "net_profit_previous": 12000,
        "operating_cashflow_current": 3000, "operating_cashflow_previous": 5000,
        "total_assets_current": 500000, "total_liabilities_current": 350000,
        "current_assets_current": 200000, "current_liabilities_current": 150000,
        "accounts_receivable_current": 40000, "accounts_receivable_previous": 20000,
        "inventory_current": 30000, "inventory_previous": 25000,
        "cost_of_goods_current": 90000, "goodwill_current": 20000,
        "net_assets_current": 150000,
    }
    data.update(overrides)
    return data


def _calc(**overrides) -> dict:
    raw = calculate_financial_indicators.invoke(
        {"financial_data_json": json.dumps(_data(**overrides), ensure_ascii=False)})
    return json.loads(raw)


def _metric(fin: dict, metric_id: str) -> dict:
    return next(item for item in fin["metric_results"] if item["metric_id"] == metric_id)


class TestSubstitution:
    """代入过程：把输入值与单位带进公式，含结果，可人工复算。"""

    def test_values_and_unit_substituted(self):
        fin = _calc()
        text = _metric(fin, "revenue_yoy_change_pct")["substitution"]
        assert "120,000.00万元" in text and "100,000.00万元" in text
        assert text.endswith("= 20.00%")

    def test_ratio_without_unit_suffix(self):
        fin = _calc()
        text = _metric(fin, "current_ratio")["substitution"]
        assert text == "200,000.00万元/150,000.00万元 = 1.33"

    def test_metric_reference_resolved_from_prior_metric(self):
        """公式引用前序指标（周转天数引用周转率）时以已算出的值为准。"""
        fin = _calc()
        text = _metric(fin, "accounts_receivable_turnover_days")["substitution"]
        assert "4.00次" in text
        assert text.endswith("= 91.25天")

    def test_rule_default_day_count_annotated(self):
        """期间天数未提供时按规则默认 365 天，须标注为规则默认而非披露值。"""
        fin = _calc()
        assert "365.00天（规则默认）" in _metric(
            fin, "accounts_receivable_turnover_days")["substitution"]

    def test_declared_day_count_used_when_provided(self):
        fin = _calc(period_days=180)
        assert "180.00天" in _metric(fin, "inventory_turnover_days")["substitution"]

    def test_missing_input_reported_not_filled(self):
        fin = _calc()
        text = _metric(fin, "other_receivables_to_total_assets_ratio_pct")["substitution"]
        assert text.startswith("输入缺失（其他应收款）")
        assert "=" not in text

    def test_zero_denominator_marked_not_computed(self):
        """总资产为零：算式照实代入，结果标未计算并带原因，不写 0。"""
        fin = _calc(total_assets_current=0)
        metric = _metric(fin, "debt_to_asset_ratio_pct")
        assert metric["value"] is None
        assert "= 未计算" in metric["substitution"]
        assert "为零" in metric["substitution"]

    def test_decimal_precision_not_pre_rounded(self):
        """代入过程按输入原值展示，不提前舍入（计算与展示精度分离）。"""
        fin = _calc(revenue_current=100000.126, revenue_previous=100000)
        text = _metric(fin, "revenue_yoy_change_pct")["substitution"]
        assert "100,000.13万元" in text


class TestIndicatorView:
    """财务指标与审计关注分析视图：分组、未计算清单与定位。"""

    def test_groups_follow_planned_order(self):
        view = build_indicator_view(_calc())
        labels = [group["label"] for group in view["groups"]]
        assert labels[:5] == ["盈利能力", "营运能力", "偿债能力", "成长能力", "现金流质量"]
        assert "资产质量与审计关注" in labels

    def test_each_metric_carries_full_chain(self):
        view = build_indicator_view(_calc())
        metric = next(m for g in view["groups"] for m in g["metrics"]
                      if m["metric_id"] == "gross_margin_pct")
        for key in ("formula", "inputs_text", "period_scope", "substitution",
                    "display_value", "status", "source_ref"):
            assert metric[key], key
        assert metric["period_scope"] == "2025年度 · 合并"
        assert "营业收入(本期)=120000万元" in metric["inputs_text"]

    def test_uncalculated_metrics_listed_with_reason(self):
        view = build_indicator_view(_calc())
        ids = {item["metric_id"]: item for item in view["uncalculated"]}
        assert "fixed_asset_turnover_ratio" in ids
        assert ids["fixed_asset_turnover_ratio"]["reason"]
        assert all(m["metric_id"] != "fixed_asset_turnover_ratio"
                   for g in view["groups"] for m in g["metrics"])

    def test_unknown_metric_still_visible(self):
        """已算出的指标必须全部出现在分组中，不得因未登记分组而静默遗漏。"""
        view = build_indicator_view(_calc())
        grouped = {m["metric_id"] for g in view["groups"] for m in g["metrics"]}
        calc = {item["metric_id"] for item in _calc()["metric_results"]
                if item["value"] is not None}
        assert calc - grouped == set()

    def test_empty_input_marks_unavailable(self):
        view = build_indicator_view("")
        assert view["available"] is False and not view["groups"]
        assert view["note"]

    def test_error_payload_marks_unavailable(self):
        assert build_indicator_view('{"error": "JSON解析失败"}')["available"] is False


class TestSourceReference:
    """原文定位：证据目录缺页码/定位时如实标注，不留空白。"""

    def test_missing_locator_stated(self):
        view = build_indicator_view(_calc())
        metric = view["groups"][0]["metrics"][0]
        assert "未提供原文定位" in metric["source_ref"]

    def test_locator_from_evidence(self):
        fin = _calc()
        for item in fin["evidence"]:
            item["page"] = "42"
            item["locator"] = "合并资产负债表"
        view = build_indicator_view(fin)
        assert "页码 42" in view["groups"][0]["metrics"][0]["source_ref"]


class TestExportSurfaces:
    """Excel 指标计算表与 PDF 财务报告须呈现同一份计算过程。"""

    def _report(self) -> dict:
        return {"company_info": {"company_name": "指标过程测试公司", "report_year": "2025"},
                "risk_details": [{"risk_id": "R001", "dimension": "财务错报风险",
                                  "title": "应收账款增速显著高于营业收入增速",
                                  "level": "重要", "evidence": "应收账款同比+100%。",
                                  "audit_suggestion": "扩大函证范围"}]}

    def test_excel_metric_sheet_has_substitution_and_locator(self, tmp_path):
        from openpyxl import load_workbook

        from tools.excel_export import _export_excel_impl

        out = tmp_path / "r.xlsx"
        _export_excel_impl(json.dumps(self._report(), ensure_ascii=False), str(out),
                           financial_indicators_json=json.dumps(_calc(), ensure_ascii=False))
        ws = load_workbook(str(out))["指标计算"]
        headers = [c.value for c in ws[1]]
        assert headers[headers.index("公式") + 1:headers.index("公式") + 3] == ["输入值及单位", "期间"]
        assert "代入过程" in headers and "原文定位" in headers
        assert headers.index("代入过程") > headers.index("口径")
        assert headers.index("代入过程") < headers.index("结果")
        rows = {str(r[0]): r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
        assert "=" in str(rows["gross_margin_pct"][headers.index("代入过程")])

    def test_pdf_financial_report_contains_chain(self, monkeypatch):
        import pypdf

        from tools.pdf_export import _export_pdf_impl

        saved = []

        def _fake(path, dest, mime):
            saved.append(path)
            return f"/local_storage/{dest}"

        monkeypatch.setattr("local_storage.upload_file_to_storage", _fake)
        monkeypatch.setattr("tools.pdf_export._embed_chart", lambda *a, **k: None)
        result = _export_pdf_impl(json.dumps(self._report(), ensure_ascii=False),
                                  financial_indicators_json=json.dumps(_calc(), ensure_ascii=False))
        assert result.count("已生成") == 3
        path = [p for p in saved if "_财务健康诊断报告_" in p][0]
        text = "".join((pg.extract_text() or "") for pg in pypdf.PdfReader(path).pages)
        text = "".join(text.split())
        assert "（三）指标计算过程与依据" in text
        assert "公式、输入值及单位、代入过程" in text
        assert "120,000.00万元-100,000.00万元" in text
        # 未计算指标不进分组表，另列名称、状态与原因，不以中性值补齐
        assert "未计算指标" in text
        assert "其他应收款/总资产" in text
