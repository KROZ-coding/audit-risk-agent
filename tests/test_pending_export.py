"""Candidates remain visible without becoming accepted risk findings."""
import json
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader
import pytest


def test_pending_findings_and_score_state_survive_all_exports(tmp_path, monkeypatch):
    from tools import pdf_export as pdf
    from tools.excel_export import _export_excel_impl

    uploaded = {}

    def upload(local_path=None, file_name=None, content_type=None, **kwargs):
        uploaded[file_name] = Path(local_path)
        return '/local_storage/' + file_name

    monkeypatch.setattr('local_storage.upload_file_to_storage', upload)
    monkeypatch.setattr(pdf, '_embed_chart', lambda *args, **kwargs: None)
    candidates = [
        {'risk_id':'R001','dimension':'财务错报风险','title':'应收回款待核查',
         'level':'重要','evidence':'经营现金流增长不排除回款关注','audit_suggestion':'核对期后回款（待执行）'},
        {'risk_id':'R002','dimension':'信息披露合规风险','title':'担保范围待核查',
         'level':'一般','evidence':'已披露不存在关联担保','audit_suggestion':'核对公告索引（待执行）'},
    ]
    for r in candidates:
        r.update(formal_status='pending',verification_status='待人工复核',
                 level_status='provisional',suggested_level=r['level'],pending_reason='外部证据未取得')
    score={'score':6.9,'base_score':6.9,'level':'低风险','level_key':'low','level_floor_adjustment':0,
           'score_status':'provisional','score_note':'量化筛查暂定，不能据此认定无风险',
           'breakdown':{'financial':0,'disclosure':23.1,'validation':0}}
    report={'company_info':{'company_name':'状态回归公司','report_year':'2025半年度'},
            'risk_details':candidates,'accepted_risk_details':[],'pending_items':candidates,
            'review_gate':{'status':'not_passed'},'comprehensive_score':score}
    out=pdf._export_split_reports(report,'','',json.dumps(score,ensure_ascii=False))
    assert out.count('已生成')==3
    texts={name:''.join(p.extract_text() or '' for p in PdfReader(path).pages)
           for name,path in uploaded.items()}
    financial=next(t for k,t in texts.items() if k.endswith('_财务健康诊断报告.pdf'))
    compliance=next(t for k,t in texts.items() if k.endswith('_合规与信息披露报告.pdf'))
    synthesis=next(t for k,t in texts.items() if k.endswith('_综合汇总报告.pdf'))
    assert '待复核提示' in financial and '待复核提示' in compliance
    assert '核对期后回款' not in financial and '核对公告索引' not in compliance
    for text in texts.values():
        assert '待复核提示' in text
        assert '不构成注册会计师审计意见' in text
    assert score['score_note'] in synthesis
    assert '系统采信相关风险 0 项' in financial
    path=tmp_path/'pending.xlsx'
    _export_excel_impl(json.dumps(report,ensure_ascii=False),str(path),comprehensive_score_json=json.dumps(score,ensure_ascii=False))
    wb=load_workbook(path)
    rows=list(wb['待复核事项'].iter_rows(min_row=2,values_only=True))
    assert len(rows)==2 and all('暂定关注' in r[3] for r in rows)
    summary={r[0]:r[1] for r in wb['报告概览'].iter_rows(values_only=True) if r[0]}
    assert summary['系统采信风险']==0 and summary['待复核提示']==2
    rules={r[0]:r[1] for r in wb['规则口径'].iter_rows(values_only=True) if r[0]}
    assert rules['基础量化分']==6.9 and rules['正式风险底线调整']==0


def test_pdf_donut_excludes_pending_counts(tmp_path, monkeypatch):
    from matplotlib.axes import Axes
    from tools.pdf_export import _gen_level_donut
    captured=[]
    original=Axes.text
    def record(self,x,y,text,*args,**kwargs):
        captured.append(str(text))
        return original(self,x,y,text,*args,**kwargs)
    monkeypatch.setattr(Axes,'text',record)
    report={'accepted_risk_details':[], 'risk_details':[
        {'risk_id':'R001','level':'重要','formal_status':'pending'},
        {'risk_id':'R002','level':'一般','formal_status':'pending'}]}
    _gen_level_donut(json.dumps(report),str(tmp_path/'donut.png'))
    assert any('系统采信0项' in s for s in captured)
    assert '2' not in captured


def test_pending_body_is_compact_review_notice():
    from reportlab.platypus import Table
    from tools.pdf_export import _pending_items_body,_build_styles,_register_chinese_font
    body=_pending_items_body({'pending_items':[{'risk_id':'R001','level':'重要',
        'pending_reason':'尚未取得完整的外部核查证据。'*25}]},_build_styles(_register_chinese_font()))
    assert not any(isinstance(item, Table) for item in body)
    text = " ".join(item.getPlainText() for item in body if hasattr(item, "getPlainText"))
    assert "1 项待复核提示" in text
    assert "未计入正式风险评分" in text


def test_pdf_review_text_hides_machine_field_names():
    from tools.pdf_export import _build_styles, _register_chinese_font, _review_conclusion_body

    report = {
        "review_conclusion": (
            "建议复核 overall_assessment、reasoning_chain 与 audit_suggestion；"
            "check_disclosure_compliance issues 应回到来源证据核对。"
        )
    }
    body = _review_conclusion_body(report, _build_styles(_register_chinese_font()))
    text = " ".join(item.getPlainText() for item in body if hasattr(item, "getPlainText"))
    for key in ("overall_assessment", "reasoning_chain", "audit_suggestion",
                "check_disclosure_compliance", " issues"):
        assert key not in text
    assert "综合评估结论" in text and "审计判断依据" in text
